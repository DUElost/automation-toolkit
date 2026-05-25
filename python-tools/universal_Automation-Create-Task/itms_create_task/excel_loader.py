# -*- coding: utf-8 -*-
"""
Excel/CSV 表单数据读取。
"""

import csv
import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook


FIELD_ALIASES = {
    "项目名": "project_name",
    "项目名称": "project_name",
    "project": "project_name",
    "project_name": "project_name",
    "主计划": "main_plan_name",
    "主计划名": "main_plan_name",
    "主计划名称": "main_plan_name",
    "main_plan_name": "main_plan_name",
    "组件": "testsuite_component_name",
    "测试组件": "testsuite_component_name",
    "模块": "testsuite_component_name",
    "testsuite_component_name": "testsuite_component_name",
    "任务类型": "task_type",
    "task_type": "task_type",
    "任务子类型": "task_sub_type",
    "测试子类型": "task_sub_type",
    "task_sub_type": "task_sub_type",
    "子计划名称": "sub_plan_name",
    "子计划名": "sub_plan_name",
    "sub_plan_name": "sub_plan_name",
    "owner": "owner_name",
    "owner_name": "owner_name",
    "负责人": "owner_name",
    "开始日期": "start_date",
    "开始时间": "start_time",
    "start_date": "start_date",
    "start_time": "start_time",
    "结束日期": "end_date",
    "结束时间": "end_time",
    "end_date": "end_date",
    "end_time": "end_time",
    "备注": "remark",
    "remark": "remark",
    "项目版本": "project_version",
    "project_version": "project_version",
    "样机标签": "project_label",
    "标签": "project_label",
    "project_label": "project_label",
    "样机阶段": "prototype_stage",
    "prototype_stage": "prototype_stage",
    "样机数量": "prototype_num",
    "prototype_num": "prototype_num",
}


def load_form_rows(file_path: str, sheet_name: Optional[str] = None, header_row: int = 1) -> List[Dict[str, Any]]:
    """读取 Excel 或 CSV 中的多行表单数据。"""
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix in (".xlsx", ".xlsm"):
        return _load_excel_rows(path, sheet_name=sheet_name, header_row=header_row)
    if suffix == ".csv":
        return _load_csv_rows(path)

    raise ValueError("暂不支持的文件类型: %s" % suffix)


def build_manpower_csv_from_weekly_plan(
    source_path: str,
    output_dir: str,
    today: Optional[date] = None,
    now: Optional[datetime] = None,
) -> str:
    """将周计划 Excel 转为人力预估 CSV。"""
    path = Path(source_path)
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("build-manpower-csv 只支持 Excel 文件")

    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel 中没有可用数据")

    headers = [_normalize_header(item) for item in rows[0]]
    required_headers = ["项目", "本周计划", "人力投入", "项目编号", "测试人员"]
    missing = [header for header in required_headers if header not in headers]
    if missing:
        raise ValueError("缺少必要表头: %s" % ",".join(missing))

    header_index = {header: index for index, header in enumerate(headers) if header}
    current_day = today or date.today()
    current_time = now or datetime.now()
    week_begin = _resolve_week_begin(path, current_day)
    week_end = week_begin + timedelta(days=6)
    task_name_suffix = _build_weekly_task_name_suffix(week_begin)
    output_path = Path(output_dir) / ("manpower_form_template_%s.csv" % current_time.strftime("%Y%m%d_%H%M%S"))
    output_path.parent.mkdir(parents=True, exist_ok=True)

    fieldnames = ["projectName", "taskName", "site", "groupType", "beginDate", "endDate", "standardWorkHours", "testerNames"]
    output_rows: List[Dict[str, Any]] = []
    for row in rows[1:]:
        project = _normalize_value(row[header_index["项目"]]) if header_index["项目"] < len(row) else None
        weekly_plan = _normalize_value(row[header_index["本周计划"]]) if header_index["本周计划"] < len(row) else None
        effort = _normalize_value(row[header_index["人力投入"]]) if header_index["人力投入"] < len(row) else None
        project_code = _normalize_value(row[header_index["项目编号"]]) if header_index["项目编号"] < len(row) else None
        tester_names = _normalize_value(row[header_index["测试人员"]]) if header_index["测试人员"] < len(row) else None
        if all(value in (None, "") for value in (project, weekly_plan, effort, project_code, tester_names)):
            continue
        if not _is_complete_weekly_plan_row(project, weekly_plan, effort, project_code):
            continue
        output_rows.append(
            {
                "projectName": project_code,
                "taskName": _build_weekly_task_name(project, weekly_plan, task_name_suffix),
                "site": "南昌",
                "groupType": "系统",
                "beginDate": week_begin.strftime("%Y-%m-%d"),
                "endDate": week_end.strftime("%Y-%m-%d"),
                "standardWorkHours": _convert_effort_to_hours(effort),
                "testerNames": tester_names or "",
            }
        )

    with open(output_path, "w", encoding="utf-8-sig", newline="") as file_obj:
        writer = csv.DictWriter(file_obj, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(output_rows)
    return str(output_path)


def load_weekly_plan_actual_rows(source_path: str, today: Optional[date] = None) -> List[Dict[str, Any]]:
    """读取原始周计划 Excel 中的测试人员与实际工时列。"""
    path = Path(source_path)
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("实际工时导入只支持 Excel 文件")

    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel 中没有可用数据")

    headers = [str(item).strip() if item is not None else None for item in rows[0]]
    required_headers = ["项目", "本周计划", "人力投入", "项目编号", "测试人员"]
    missing = [header for header in required_headers if header not in headers]
    if missing:
        raise ValueError("缺少必要表头: %s" % ",".join(missing))

    actual_headers = [header for header in headers if header and header.startswith("实际工时")]
    actual_headers.sort(key=_sort_actual_header_key)
    header_index = {header: index for index, header in enumerate(headers) if header}

    current_day = today or date.today()
    week_begin = _resolve_week_begin(path, current_day)
    week_end = week_begin + timedelta(days=6)
    task_name_suffix = _build_weekly_task_name_suffix(week_begin)

    output_rows: List[Dict[str, Any]] = []
    for row in rows[1:]:
        project = _normalize_value(row[header_index["项目"]]) if header_index["项目"] < len(row) else None
        weekly_plan = _normalize_value(row[header_index["本周计划"]]) if header_index["本周计划"] < len(row) else None
        effort = _normalize_value(row[header_index["人力投入"]]) if header_index["人力投入"] < len(row) else None
        project_code = _normalize_value(row[header_index["项目编号"]]) if header_index["项目编号"] < len(row) else None
        tester_names = _normalize_value(row[header_index["测试人员"]]) if header_index["测试人员"] < len(row) else None
        if all(value in (None, "") for value in (project, weekly_plan, effort, project_code, tester_names)):
            continue
        if not _is_complete_weekly_plan_row(project, weekly_plan, effort, project_code):
            continue

        output_rows.append(
            {
                "projectName": project_code,
                "taskName": _build_weekly_task_name(project, weekly_plan, task_name_suffix),
                "beginDate": week_begin.strftime("%Y-%m-%d"),
                "endDate": week_end.strftime("%Y-%m-%d"),
                "effort": effort,
                "testerNames": _split_comma_values(tester_names),
                "actualRatios": [
                    _normalize_value(row[header_index[header]]) if header_index[header] < len(row) else None
                    for header in actual_headers
                ],
            }
        )
    return output_rows


def _load_excel_rows(path: Path, sheet_name: Optional[str], header_row: int) -> List[Dict[str, Any]]:
    """读取 Excel。"""
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ValueError("未找到工作表: %s" % sheet_name)
        sheet = workbook[sheet_name]
    else:
        sheet = workbook[workbook.sheetnames[0]]

    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < header_row:
        return []

    headers = [_normalize_header(item) for item in rows[header_row - 1]]
    results: List[Dict[str, Any]] = []

    for row in rows[header_row:]:
        item: Dict[str, Any] = {}
        has_value = False
        for index, header in enumerate(headers):
            if not header:
                continue
            value = row[index] if index < len(row) else None
            normalized = _normalize_value(value)
            if normalized not in (None, ""):
                has_value = True
            item[header] = normalized
        if has_value:
            results.append(_drop_empty_values(item))

    return results


def _load_csv_rows(path: Path) -> List[Dict[str, Any]]:
    """读取 CSV。"""
    results: List[Dict[str, Any]] = []
    with open(path, "r", encoding="utf-8-sig", newline="") as file_obj:
        reader = csv.DictReader(file_obj)
        for row in reader:
            item = {}
            has_value = False
            for key, value in row.items():
                header = _normalize_header(key)
                if not header:
                    continue
                normalized = _normalize_value(value)
                if normalized not in (None, ""):
                    has_value = True
                item[header] = normalized
            if has_value:
                results.append(_drop_empty_values(item))
    return results


def _normalize_header(value: Any) -> Optional[str]:
    """标准化表头。"""
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    return FIELD_ALIASES.get(text, FIELD_ALIASES.get(text.lower(), text))


def _normalize_value(value: Any) -> Any:
    """标准化单元格值。"""
    if value is None:
        return None
    if isinstance(value, datetime):
        if value.hour == 0 and value.minute == 0 and value.second == 0:
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return value

    text = str(value).strip()
    if not text:
        return None

    if text.startswith("{") or text.startswith("["):
        try:
            return json.loads(text)
        except json.JSONDecodeError:
            return text
    return text


def _drop_empty_values(item: Dict[str, Any]) -> Dict[str, Any]:
    """删除空值字段。"""
    return {key: value for key, value in item.items() if value not in (None, "")}


def _build_weekly_task_name(project: Any, weekly_plan: Any, week_suffix: str) -> str:
    """构造附带周次的人力预估任务名称。"""
    return "%s：%s%s" % (project, weekly_plan, week_suffix)


def _resolve_week_begin(path: Path, current_day: date) -> date:
    week_number = _extract_week_number_from_filename(path)
    if week_number is None:
        return current_day - timedelta(days=current_day.weekday())
    return _closest_iso_week_begin(week_number, current_day)


def _extract_week_number_from_filename(path: Path) -> Optional[int]:
    match = re.search(r"(?i)WK\s*(\d{1,2})", path.stem)
    if not match:
        return None
    return int(match.group(1))


def _closest_iso_week_begin(week_number: int, current_day: date) -> date:
    current_iso_year = current_day.isocalendar()[0]
    candidates: List[date] = []
    for iso_year in range(current_iso_year - 1, current_iso_year + 2):
        try:
            candidates.append(date.fromisocalendar(iso_year, week_number, 1))
        except ValueError:
            continue
    if not candidates:
        raise ValueError("文件名中的 WK 周次不合法: WK%s" % week_number)
    return min(candidates, key=lambda item: abs((item - current_day).days))


def _build_weekly_task_name_suffix(week_begin: date) -> str:
    """根据周起始日期生成任务名周次后缀。"""
    _, iso_week, _ = week_begin.isocalendar()
    return "（WK%s）" % iso_week


def _is_complete_weekly_plan_row(project: Any, weekly_plan: Any, effort: Any, project_code: Any) -> bool:
    """判断周计划行是否具备构造任务的核心字段。"""
    return all(value not in (None, "") for value in (project, weekly_plan, effort, project_code))


def _convert_effort_to_hours(value: Any) -> Any:
    """将人力投入转换为标准工时。"""
    if value in (None, ""):
        raise ValueError("人力投入不能为空")
    if isinstance(value, bool):
        raise ValueError("人力投入必须是数字")
    if isinstance(value, (int, float)):
        result = value * 8
    else:
        try:
            result = float(str(value))
        except ValueError as exc:
            raise ValueError("人力投入必须是数字") from exc
        result *= 8
    if isinstance(result, float) and result.is_integer():
        return int(result)
    return result


def _split_comma_values(value: Any) -> List[str]:
    """将逗号分隔字符串切成列表。"""
    results: List[str] = []
    for part in str(value or "").split(","):
        normalized = str(part).strip()
        if normalized:
            results.append(normalized)
    return results


def _sort_actual_header_key(header: str) -> int:
    """按实际工时列编号排序。"""
    digits = "".join(char for char in header if char.isdigit())
    if digits:
        return int(digits)
    return 0
