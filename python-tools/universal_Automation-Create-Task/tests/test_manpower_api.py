# -*- coding: utf-8 -*-
import csv
import json
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from itms_create_task.config_loader import load_manpower_config
from itms_create_task.excel_loader import build_manpower_csv_from_weekly_plan, load_form_rows, load_weekly_plan_actual_rows
import itms_create_task.workflow as workflow
from itms_create_task.manpower_api import ManpowerApiClient

MANPOWER_CONFIG_PATH = "examples/manpower_config.example.json"
MANPOWER_FORM_PATH = "examples/manpower_form.example.json"


class DummyResponse:
    def __init__(self, payload=None, ok=True, json_error=None):
        self._payload = payload
        self.ok = ok
        self._json_error = json_error

    def json(self):
        if self._json_error is not None:
            raise self._json_error
        return self._payload


class DummySession:
    def __init__(self, response):
        self.response = response
        self.calls = []

    def get(self, url, params=None, timeout=None):
        self.calls.append({"url": url, "params": params, "timeout": timeout})
        return self.response


class DummyAuthManager:
    def __init__(self, response):
        self.session = DummySession(response)

    def create_requests_session(self, auto_refresh=True):
        return self.session


class RecordingResponse:
    def __init__(self, payload=None, text="", status_code=200, ok=True, json_error=None):
        self._payload = payload
        self.text = text
        self.status_code = status_code
        self.ok = ok
        self._json_error = json_error

    def json(self):
        if self._json_error is not None:
            raise self._json_error
        return self._payload


class RecordingSession(DummySession):
    def __init__(self, response):
        super().__init__(response)
        self.requests = []

    def request(self, method, url, json=None, headers=None, timeout=None):
        self.requests.append(
            {
                "method": method,
                "url": url,
                "json": json,
                "headers": headers,
                "timeout": timeout,
            }
        )
        return self.response

    def post(self, url, json=None, headers=None, timeout=None):
        return self.request("POST", url, json=json, headers=headers, timeout=timeout)


class RecordingAuthManager:
    def __init__(self, response):
        self.session = RecordingSession(response)

    def create_requests_session(self, auto_refresh=True):
        return self.session


class CreateEstimateSession(DummySession):
    def __init__(self, auth_response, post_response=None, post_error=None):
        super().__init__(auth_response)
        self.post_response = post_response
        self.post_error = post_error
        self.requests = []

    def post(self, url, json=None, headers=None, timeout=None):
        self.requests.append(
            {
                "method": "POST",
                "url": url,
                "json": json,
                "headers": headers,
                "timeout": timeout,
            }
        )
        if self.post_error is not None:
            raise self.post_error
        return self.post_response


class CreateEstimateAuthManager:
    def __init__(self, auth_response, post_response=None, post_error=None):
        self.session = CreateEstimateSession(auth_response, post_response=post_response, post_error=post_error)

    def create_requests_session(self, auto_refresh=True):
        return self.session


class SequenceSession:
    def __init__(self, get_responses=None, post_responses=None, post_error=None):
        self.get_responses = list(get_responses or [])
        self.post_responses = list(post_responses or [])
        self.post_error = post_error
        self.get_calls = []
        self.post_calls = []

    def get(self, url, params=None, timeout=None):
        self.get_calls.append({"url": url, "params": params, "timeout": timeout})
        if self.get_responses:
            return self.get_responses.pop(0)
        raise AssertionError("未准备 GET 响应")

    def post(self, url, json=None, data=None, headers=None, timeout=None):
        self.post_calls.append({"url": url, "json": json, "data": data, "headers": headers, "timeout": timeout})
        if self.post_error is not None:
            raise self.post_error
        if self.post_responses:
            return self.post_responses.pop(0)
        raise AssertionError("未准备 POST 响应")


class SequenceAuthManager:
    def __init__(self, get_responses=None, post_responses=None, post_error=None):
        self.session = SequenceSession(get_responses=get_responses, post_responses=post_responses, post_error=post_error)

    def create_requests_session(self, auto_refresh=True):
        return self.session


def build_config():
    return SimpleNamespace(base_url="http://itms.tinno.com")


def test_build_request_payload_uses_current_user_id():
    auth_manager = DummyAuthManager(DummyResponse({"errno": 0, "data": {"userId": "65417411391957279", "nickName": "桂雨旺"}}))
    client = ManpowerApiClient(build_config(), auth_manager)

    payload = client.build_request_payload(
        {
            "projectName": "V657",
            "taskName": "第九轮稳定性挂测",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026-04-06",
            "endDate": "2026-04-12",
            "standardWorkHours": 32,
        }
    )

    assert payload["leaderId"] == "65417411391957279"


def test_build_request_payload_generates_single_week_daterange():
    auth_manager = DummyAuthManager(DummyResponse({"errno": 0, "data": {"id": "65417411391957279"}}))
    client = ManpowerApiClient(build_config(), auth_manager)

    payload = client.build_request_payload(
        {
            "projectName": "V657",
            "taskName": "第九轮稳定性挂测",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026-04-06",
            "endDate": "2026-04-12",
            "standardWorkHours": 32,
        }
    )

    assert payload["daterange"] == ["2026-04-05T16:00:00.000Z", "2026-04-11T16:00:00.000Z"]


def test_build_request_payload_generates_single_week_list():
    auth_manager = DummyAuthManager(DummyResponse({"errno": 0, "data": {"userId": "65417411391957279"}}))
    client = ManpowerApiClient(build_config(), auth_manager)

    payload = client.build_request_payload(
        {
            "projectName": "V657",
            "taskName": "第九轮稳定性挂测",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026-04-06",
            "endDate": "2026-04-12",
            "standardWorkHours": 32,
        }
    )

    assert payload["weekList"] == [
        {
            "year": "2026",
            "week": 15,
            "splitBeginDate": "2026-04-06",
            "splitEndDate": "2026-04-12",
            "standardWorkHours": 32,
        }
    ]


def test_get_current_user_raises_when_errno_is_not_zero():
    auth_manager = DummyAuthManager(DummyResponse({"errno": 401, "errmsg": "token expired", "data": {}}))
    client = ManpowerApiClient(build_config(), auth_manager)

    with pytest.raises(RuntimeError, match="token expired|errno"):
        client.get_current_user()


def test_get_current_user_raises_when_response_is_not_ok():
    auth_manager = DummyAuthManager(DummyResponse({"errno": 0, "data": {}}, ok=False))
    client = ManpowerApiClient(build_config(), auth_manager)

    with pytest.raises(RuntimeError, match="auth/info"):
        client.get_current_user()


def test_get_current_user_raises_when_json_decode_fails():
    auth_manager = DummyAuthManager(DummyResponse(json_error=ValueError("bad json")))
    client = ManpowerApiClient(build_config(), auth_manager)

    with pytest.raises(RuntimeError, match="auth/info|bad json"):
        client.get_current_user()


def test_get_current_user_raises_when_data_is_not_object():
    auth_manager = DummyAuthManager(DummyResponse({"errno": 0, "data": ["not", "dict"]}))
    client = ManpowerApiClient(build_config(), auth_manager)

    with pytest.raises(RuntimeError, match="data"):
        client.get_current_user()


def test_build_request_payload_rejects_invalid_date_format():
    auth_manager = DummyAuthManager(DummyResponse({"errno": 0, "data": {"id": "65417411391957279"}}))
    client = ManpowerApiClient(build_config(), auth_manager)

    with pytest.raises(ValueError, match="YYYY-MM-DD"):
        client.build_request_payload(
            {
                "projectName": "V657",
                "taskName": "日期非法",
                "site": "南昌",
                "groupType": "系统",
                "beginDate": "2026-04-06 00:00:00",
                "endDate": "2026-04-12",
                "standardWorkHours": 32,
            }
        )


def test_build_request_payload_accepts_slash_date_format():
    auth_manager = DummyAuthManager(DummyResponse({"errno": 0, "data": {"id": "65417411391957279"}}))
    client = ManpowerApiClient(build_config(), auth_manager)

    payload = client.build_request_payload(
        {
            "projectName": "V657",
            "taskName": "斜杠日期",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026/4/6",
            "endDate": "2026/4/12",
            "standardWorkHours": 32,
        }
    )

    assert payload["beginDate"] == "2026-04-06"
    assert payload["endDate"] == "2026-04-12"
    assert payload["daterange"] == ["2026-04-05T16:00:00.000Z", "2026-04-11T16:00:00.000Z"]


def test_build_request_payload_rejects_begin_date_after_end_date():
    auth_manager = DummyAuthManager(DummyResponse({"errno": 0, "data": {"id": "65417411391957279"}}))
    client = ManpowerApiClient(build_config(), auth_manager)

    with pytest.raises(ValueError, match="beginDate"):
        client.build_request_payload(
            {
                "projectName": "V657",
                "taskName": "日期反转",
                "site": "南昌",
                "groupType": "系统",
                "beginDate": "2026-04-13",
                "endDate": "2026-04-12",
                "standardWorkHours": 32,
            }
        )


def test_build_request_payload_splits_cross_week_ranges():
    auth_manager = DummyAuthManager(DummyResponse({"errno": 0, "data": {"userId": "65417411391957279"}}))
    client = ManpowerApiClient(build_config(), auth_manager)

    payload = client.build_request_payload(
        {
            "projectName": "V657",
            "taskName": "跨周任务",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026-04-10",
            "endDate": "2026-04-14",
            "standardWorkHours": "32",
        }
    )

    assert payload["weekList"] == [
        {
            "year": "2026",
            "week": 15,
            "splitBeginDate": "2026-04-10",
            "splitEndDate": "2026-04-12",
            "standardWorkHours": 32,
        },
        {
            "year": "2026",
            "week": 16,
            "splitBeginDate": "2026-04-13",
            "splitEndDate": "2026-04-14",
            "standardWorkHours": 32,
        },
    ]
    assert payload["standardWorkHours"] == 32


def test_build_request_payload_rejects_invalid_standard_work_hours():
    auth_manager = DummyAuthManager(DummyResponse({"errno": 0, "data": {"userId": "65417411391957279"}}))
    client = ManpowerApiClient(build_config(), auth_manager)

    with pytest.raises(ValueError, match="standardWorkHours"):
        client.build_request_payload(
            {
                "projectName": "V657",
                "taskName": "工时非法",
                "site": "南昌",
                "groupType": "系统",
                "beginDate": "2026-04-10",
                "endDate": "2026-04-14",
                "standardWorkHours": "3.5",
            }
        )


def test_build_request_payload_rejects_missing_standard_work_hours():
    auth_manager = DummyAuthManager(DummyResponse({"errno": 0, "data": {"userId": "65417411391957279"}}))
    client = ManpowerApiClient(build_config(), auth_manager)

    with pytest.raises(ValueError, match="standardWorkHours"):
        client.build_request_payload(
            {
                "projectName": "V657",
                "taskName": "缺少工时",
                "site": "南昌",
                "groupType": "系统",
                "beginDate": "2026-04-10",
                "endDate": "2026-04-14",
            }
        )


def test_build_request_payload_rejects_blank_required_fields():
    auth_manager = DummyAuthManager(DummyResponse({"errno": 0, "data": {"userId": "65417411391957279"}}))
    client = ManpowerApiClient(build_config(), auth_manager)

    with pytest.raises(ValueError, match="projectName|taskName"):
        client.build_request_payload(
            {
                "projectName": "   ",
                "taskName": "\t",
                "site": "南昌",
                "groupType": "系统",
                "beginDate": "2026-04-10",
                "endDate": "2026-04-14",
                "standardWorkHours": 32,
            }
        )


@pytest.mark.parametrize(
    "field_name, field_value",
    [
        ("projectName", 0),
        ("taskName", False),
        ("site", []),
    ],
)
def test_build_request_payload_rejects_non_string_text_fields(field_name, field_value):
    auth_manager = DummyAuthManager(DummyResponse({"errno": 0, "data": {"userId": "65417411391957279"}}))
    client = ManpowerApiClient(build_config(), auth_manager)

    form_data = {
        "projectName": "V657",
        "taskName": "非字符串字段",
        "site": "南昌",
        "groupType": "系统",
        "beginDate": "2026-04-10",
        "endDate": "2026-04-14",
        "standardWorkHours": 32,
    }
    form_data[field_name] = field_value

    with pytest.raises(ValueError, match=field_name):
        client.build_request_payload(form_data)


def test_create_estimate_posts_to_hrestimate_create():
    auth_manager = RecordingAuthManager(RecordingResponse({"errno": 0, "data": {"id": "ok"}}, text='{"errno":0,"data":{"id":"ok"}}'))
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.create_estimate(
        {
            "projectName": "V657",
            "taskName": "接口提交",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026-04-10",
            "endDate": "2026-04-14",
            "standardWorkHours": 32,
        }
    )

    assert auth_manager.session.requests == [
        {
            "method": "POST",
            "url": "http://itms.tinno.com/api/task/hrestimate/create",
            "json": result["request_body"],
            "headers": {"content-type": "application/json;charset=UTF-8"},
            "timeout": 30,
        }
    ]


def test_create_estimate_returns_standardized_result():
    auth_manager = RecordingAuthManager(
        RecordingResponse({"errno": 0, "data": {"id": "ok"}}, text='{"errno":0,"data":{"id":"ok"}}', status_code=200)
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.create_estimate(
        {
            "projectName": "V657",
            "taskName": "标准化结果",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026-04-10",
            "endDate": "2026-04-14",
            "standardWorkHours": 32,
        }
    )

    assert set(result) >= {"success", "status_code", "request_body", "response_body"}
    assert result["success"] is True
    assert result["status_code"] == 200
    assert isinstance(result["request_body"], dict)
    assert result["response_body"] == '{"errno":0,"data":{"id":"ok"}}'


def test_create_estimate_fails_when_response_json_is_not_dict():
    auth_manager = CreateEstimateAuthManager(
        DummyResponse({"errno": 0, "data": {"userId": "65417411391957279"}}),
        post_response=RecordingResponse(["not", "dict"], text='["not","dict"]', status_code=200),
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.create_estimate(
        {
            "projectName": "V657",
            "taskName": "非字典响应",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026-04-10",
            "endDate": "2026-04-14",
            "standardWorkHours": 32,
        }
    )

    assert result["success"] is False
    assert result["status_code"] == 200


def test_create_estimate_fails_when_response_json_decode_fails():
    auth_manager = CreateEstimateAuthManager(
        DummyResponse({"errno": 0, "data": {"userId": "65417411391957279"}}),
        post_response=RecordingResponse(payload=None, text="not json", status_code=200, json_error=ValueError("bad json")),
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.create_estimate(
        {
            "projectName": "V657",
            "taskName": "解析失败",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026-04-10",
            "endDate": "2026-04-14",
            "standardWorkHours": 32,
        }
    )

    assert result["success"] is False
    assert result["status_code"] == 200
    assert result["response_body"] == "not json"


def test_create_estimate_returns_standardized_result_when_post_raises():
    auth_manager = CreateEstimateAuthManager(
        DummyResponse({"errno": 0, "data": {"userId": "65417411391957279"}}),
        post_error=RuntimeError("network down"),
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.create_estimate(
        {
            "projectName": "V657",
            "taskName": "请求异常",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026-04-10",
            "endDate": "2026-04-14",
            "standardWorkHours": 32,
        }
    )

    assert result["success"] is False
    assert result["status_code"] is None
    assert "network down" in result["response_body"]


def test_create_estimate_returns_failure_when_build_request_payload_raises(monkeypatch):
    auth_manager = CreateEstimateAuthManager(
        DummyResponse({"errno": 0, "data": {"userId": "65417411391957279"}}),
        post_response=RecordingResponse({"errno": 0}, text='{"errno":0}', status_code=200),
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    def boom(_form_data):
        raise ValueError("payload broken")

    monkeypatch.setattr(client, "build_request_payload", boom)

    result = client.create_estimate(
        {
            "projectName": "V657",
            "taskName": "构造失败",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026-04-10",
            "endDate": "2026-04-14",
            "standardWorkHours": 32,
        }
    )

    assert result["success"] is False
    assert result["status_code"] is None
    assert result["request_body"] is None
    assert "payload broken" in result["response_body"]


def test_create_estimate_returns_failure_when_get_current_user_raises():
    auth_manager = CreateEstimateAuthManager(
        DummyResponse({"errno": 401, "errmsg": "token expired", "data": {}}),
        post_response=RecordingResponse({"errno": 0}, text='{"errno":0}', status_code=200),
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.create_estimate(
        {
            "projectName": "V657",
            "taskName": "用户获取失败",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026-04-10",
            "endDate": "2026-04-14",
            "standardWorkHours": 32,
        }
    )

    assert result["success"] is False
    assert result["status_code"] is None
    assert result["request_body"] is None
    assert "token expired" in result["response_body"]


def test_create_estimate_fails_when_response_ok_is_false():
    auth_manager = CreateEstimateAuthManager(
        DummyResponse({"errno": 0, "data": {"userId": "65417411391957279"}}),
        post_response=RecordingResponse({"errno": 0, "data": {"id": "ok"}}, text='{"errno":0,"data":{"id":"ok"}}', status_code=200, ok=False),
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.create_estimate(
        {
            "projectName": "V657",
            "taskName": "http失败",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026-04-10",
            "endDate": "2026-04-14",
            "standardWorkHours": 32,
        }
    )

    assert result["success"] is False
    assert result["status_code"] == 200


def test_create_estimate_fails_when_response_errno_is_not_zero():
    auth_manager = CreateEstimateAuthManager(
        DummyResponse({"errno": 0, "data": {"userId": "65417411391957279"}}),
        post_response=RecordingResponse({"errno": 1, "errmsg": "bad request"}, text='{"errno":1,"errmsg":"bad request"}', status_code=200),
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.create_estimate(
        {
            "projectName": "V657",
            "taskName": "业务失败",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026-04-10",
            "endDate": "2026-04-14",
            "standardWorkHours": 32,
        }
    )

    assert result["success"] is False
    assert result["status_code"] == 200


def test_create_estimate_reuses_build_request_payload(monkeypatch):
    auth_manager = RecordingAuthManager(RecordingResponse({"errno": 0, "data": {"id": "ok"}}, text='{"errno":0,"data":{"id":"ok"}}'))
    client = ManpowerApiClient(build_config(), auth_manager)
    payload = {"sentinel": True}
    calls = []

    def fake_build_request_payload(form_data):
        calls.append(form_data)
        return payload

    monkeypatch.setattr(client, "build_request_payload", fake_build_request_payload)

    result = client.create_estimate(
        {
            "projectName": "V657",
            "taskName": "复用请求体",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026-04-10",
            "endDate": "2026-04-14",
            "standardWorkHours": 32,
        }
    )

    assert calls and calls[0]["taskName"] == "复用请求体"
    assert result["request_body"] is payload
    assert auth_manager.session.requests[0]["json"] is payload


def test_resolve_tester_names_matches_username_and_nickname():
    auth_manager = SequenceAuthManager(
        get_responses=[
            RecordingResponse(
                {"errno": 0, "data": {"list": [
                    {"id": "u1", "username": "TNNC01", "nickname": "张三", "status": 1},
                    {"id": "u2", "username": "lisi", "nickname": "李四", "status": 1},
                ]}},
                text='{"errno":0}',
            )
        ]
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    resolved = client.resolve_tester_names("TNNC01,李四")

    assert resolved["success"] is True
    assert resolved["tester_ids"] == ["u1", "u2"]
    assert resolved["tester_names"] == ["TNNC01", "李四"]


def test_resolve_tester_names_fails_when_user_missing():
    auth_manager = SequenceAuthManager(
        get_responses=[RecordingResponse({"errno": 0, "data": {"list": []}}, text='{"errno":0}')]
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    resolved = client.resolve_tester_names("不存在")

    assert resolved["success"] is False
    assert "未找到测试人员" in resolved["error"]


def test_resolve_tester_names_fails_when_nickname_duplicated():
    auth_manager = SequenceAuthManager(
        get_responses=[
            RecordingResponse(
                {"errno": 0, "data": {"list": [
                    {"id": "u1", "username": "zhangsan1", "nickname": "张三", "status": 1},
                    {"id": "u2", "username": "zhangsan2", "nickname": "张三", "status": 1},
                ]}},
                text='{"errno":0}',
            )
        ]
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    resolved = client.resolve_tester_names("张三")

    assert resolved["success"] is False
    assert "匹配到多个用户" in resolved["error"]


def test_assign_testers_posts_update_tester_with_estimate_id():
    auth_manager = SequenceAuthManager(
        get_responses=[
            RecordingResponse(
                {"errno": 0, "data": {"list": [
                    {"id": "u1", "username": "TNNC01", "nickname": "张三", "status": 1}
                ]}},
                text='{"errno":0}',
            )
        ],
        post_responses=[RecordingResponse({"errno": 0, "errmsg": "成功"}, text='{"errno":0,"errmsg":"成功"}')],
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.assign_testers("estimate-1", "TNNC01")

    assert result["success"] is True
    assert auth_manager.session.post_calls[0]["json"]["id"] == "estimate-1"
    assert auth_manager.session.post_calls[0]["json"]["userIds"] == ["u1"]


def test_assign_testers_returns_failure_when_api_errno_not_zero():
    auth_manager = SequenceAuthManager(
        get_responses=[
            RecordingResponse(
                {"errno": 0, "data": {"list": [
                    {"id": "u1", "username": "TNNC01", "nickname": "张三", "status": 1}
                ]}},
                text='{"errno":0}',
            )
        ],
        post_responses=[RecordingResponse({"errno": 500, "errmsg": "失败"}, text='{"errno":500,"errmsg":"失败"}')],
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.assign_testers("estimate-1", "TNNC01")

    assert result["success"] is False
    assert result["request_body"]["id"] == "estimate-1"
    assert "失败" in result["response_body"]


def test_fill_actual_work_posts_update_actual_with_weekly_personal_list():
    auth_manager = SequenceAuthManager(
        get_responses=[
            RecordingResponse(
                {
                    "errno": 0,
                    "data": {
                        "id": "estimate-1",
                        "weekList": [
                            {
                                "id": "weekly-1",
                                "taskEstimateId": "estimate-1",
                                "year": 2026,
                                "week": 15,
                                "splitBeginDate": "2026-04-06",
                                "splitEndDate": "2026-04-12",
                                "standardWorkHours": 48,
                            }
                        ],
                    },
                },
                text='{"errno":0}',
            ),
            RecordingResponse(
                {
                    "errno": 0,
                    "data": {
                        "list": [
                            {"id": "u1", "username": "TNNC01", "nickname": "吕代", "status": 1},
                            {"id": "u2", "username": "TNNC02", "nickname": "张三", "status": 1},
                        ]
                    },
                },
                text='{"errno":0}',
            ),
        ],
        post_responses=[RecordingResponse({"errno": 0, "errmsg": "成功"}, text='{"errno":0,"errmsg":"成功"}')],
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.fill_actual_work("estimate-1", effort=8, tester_names=["吕代", "张三"], actual_ratios=[0.1, 0.2])

    assert result["success"] is True
    assert auth_manager.session.get_calls[0]["url"] == "http://itms.tinno.com/api/task/hrestimate/read"
    assert auth_manager.session.get_calls[0]["params"] == {"id": "estimate-1"}
    request_body = auth_manager.session.post_calls[0]["json"]
    assert request_body["id"] == "estimate-1"
    assert request_body["weekList"][0]["id"] == "weekly-1"
    assert request_body["weekList"][0]["weeklyPersonalList"] == [
        {
            "taskEstimateId": "estimate-1",
            "taskWeeklyId": "weekly-1",
            "userId": "u1",
            "userName": "吕代",
            "personalActualWorkHours": 6.4,
        },
        {
            "taskEstimateId": "estimate-1",
            "taskWeeklyId": "weekly-1",
            "userId": "u2",
            "userName": "张三",
            "personalActualWorkHours": 12.8,
        },
    ]


def test_fill_actual_work_reads_week_list_from_hrweekly_list_when_read_missing_weeklist():
    auth_manager = SequenceAuthManager(
        get_responses=[
            RecordingResponse(
                {
                    "errno": 0,
                    "data": {
                        "id": "estimate-1",
                        "projectName": "平台类",
                    },
                },
                text='{"errno":0}',
            ),
            RecordingResponse(
                {
                    "errno": 0,
                    "data": [
                        {
                            "id": "weekly-1",
                            "taskEstimateId": "estimate-1",
                            "year": 2026,
                            "week": 15,
                            "splitBeginDate": "2026-04-06",
                            "splitEndDate": "2026-04-12",
                            "standardWorkHours": 48,
                        }
                    ],
                },
                text='{"errno":0}',
            ),
            RecordingResponse(
                {
                    "errno": 0,
                    "data": {
                        "list": [
                            {"id": "u1", "username": "TNNC01", "nickname": "洪子鑫", "status": 1},
                        ]
                    },
                },
                text='{"errno":0}',
            ),
        ],
        post_responses=[RecordingResponse({"errno": 0, "errmsg": "成功"}, text='{"errno":0,"errmsg":"成功"}')],
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.fill_actual_work("estimate-1", effort=6, tester_names=["洪子鑫"], actual_ratios=[1])

    assert result["success"] is True
    assert auth_manager.session.get_calls[1]["url"] == "http://itms.tinno.com/api/task/hrweekly/list"
    assert auth_manager.session.get_calls[1]["params"] == {"taskEstimateId": "estimate-1"}
    assert auth_manager.session.post_calls[0]["json"]["weekList"][0]["weeklyPersonalList"][0]["personalActualWorkHours"] == 48


def test_fill_actual_work_fails_when_estimate_has_multiple_weeks():
    auth_manager = SequenceAuthManager(
        get_responses=[
            RecordingResponse(
                {
                    "errno": 0,
                    "data": {
                        "id": "estimate-1",
                        "weekList": [
                            {"id": "weekly-1", "taskEstimateId": "estimate-1", "year": 2026, "week": 15, "splitBeginDate": "2026-04-06", "splitEndDate": "2026-04-12", "standardWorkHours": 24},
                            {"id": "weekly-2", "taskEstimateId": "estimate-1", "year": 2026, "week": 16, "splitBeginDate": "2026-04-13", "splitEndDate": "2026-04-19", "standardWorkHours": 24},
                        ],
                    },
                },
                text='{"errno":0}',
            )
        ]
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.fill_actual_work("estimate-1", effort=8, tester_names=["吕代"], actual_ratios=[0.1])

    assert result["success"] is False
    assert result["response_body"] == "当前仅支持单周任务"
    assert auth_manager.session.post_calls == []


def test_complete_estimate_posts_form_encoded_status_request():
    auth_manager = SequenceAuthManager(
        post_responses=[RecordingResponse({"errno": 0, "errmsg": "成功"}, text='{"errno":0,"errmsg":"成功"}')]
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.complete_estimate("estimate-1")

    assert result["success"] is True
    assert auth_manager.session.post_calls[0]["url"] == "http://itms.tinno.com/api/task/hrestimate/update/status"
    assert auth_manager.session.post_calls[0]["data"] == "ids=estimate-1&status=2"
    assert auth_manager.session.post_calls[0]["headers"]["content-type"] == "application/x-www-form-urlencoded; charset=UTF-8"


def test_find_estimate_matches_by_project_task_and_date():
    auth_manager = SequenceAuthManager(
        get_responses=[
            RecordingResponse({"errno": 0, "data": {"userId": "leader-1"}}, text='{"errno":0}'),
            RecordingResponse(
                {"errno": 0, "data": {"list": [
                    {
                        "id": "estimate-1",
                        "projectName": "T727",
                        "taskName": "X1103：专项验证",
                        "beginDate": "2026-04-06",
                        "endDate": "2026-04-12",
                        "leaderId": "leader-1",
                    }
                ]}},
                text='{"errno":0}',
            ),
        ]
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.find_estimate(
        {
            "projectName": "T727",
            "taskName": "X1103：专项验证",
            "beginDate": "2026/4/6",
            "endDate": "2026/4/12",
        }
    )

    assert result["success"] is True
    assert result["estimate_id"] == "estimate-1"


def test_find_estimate_matches_legacy_task_name_without_week_suffix():
    auth_manager = SequenceAuthManager(
        get_responses=[
            RecordingResponse({"errno": 0, "data": {"userId": "leader-1"}}, text='{"errno":0}'),
            RecordingResponse(
                {"errno": 0, "data": {"list": [
                    {
                        "id": "estimate-1",
                        "projectName": "T727",
                        "taskName": "X1103：专项验证",
                        "beginDate": "2026-04-06",
                        "endDate": "2026-04-12",
                        "leaderId": "leader-1",
                    }
                ]}},
                text='{"errno":0}',
            ),
        ]
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.find_estimate(
        {
            "projectName": "T727",
            "taskName": "X1103：专项验证（WK15）",
            "beginDate": "2026-04-06",
            "endDate": "2026-04-12",
        }
    )

    assert result["success"] is True
    assert result["estimate_id"] == "estimate-1"


def test_find_estimate_fails_when_no_match():
    auth_manager = SequenceAuthManager(
        get_responses=[
            RecordingResponse({"errno": 0, "data": {"userId": "leader-1"}}, text='{"errno":0}'),
            RecordingResponse({"errno": 0, "data": {"list": []}}, text='{"errno":0}'),
        ]
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.find_estimate(
        {
            "projectName": "T727",
            "taskName": "X1103：专项验证",
            "beginDate": "2026-04-06",
            "endDate": "2026-04-12",
        }
    )

    assert result["success"] is False
    assert "未找到匹配的人力预估任务" in result["error"]


def test_find_estimate_fails_when_multiple_matches():
    auth_manager = SequenceAuthManager(
        get_responses=[
            RecordingResponse({"errno": 0, "data": {"userId": "leader-1"}}, text='{"errno":0}'),
            RecordingResponse(
                {"errno": 0, "data": {"list": [
                    {"id": "estimate-1", "projectName": "T727", "taskName": "X1103：专项验证", "beginDate": "2026-04-06", "endDate": "2026-04-12"},
                    {"id": "estimate-2", "projectName": "T727", "taskName": "X1103：专项验证", "beginDate": "2026-04-06", "endDate": "2026-04-12"},
                ]}},
                text='{"errno":0}',
            ),
        ]
    )
    client = ManpowerApiClient(build_config(), auth_manager)

    result = client.find_estimate(
        {
            "projectName": "T727",
            "taskName": "X1103：专项验证",
            "beginDate": "2026-04-06",
            "endDate": "2026-04-12",
        }
    )

    assert result["success"] is False
    assert "匹配到多条人力预估任务" in result["error"]


def test_run_create_then_assign_uses_created_estimate_id(monkeypatch):
    client = ManpowerApiClient(build_config(), SequenceAuthManager())
    monkeypatch.setattr(
        client,
        "create_estimate",
        lambda form_data: {
            "success": True,
            "status_code": 200,
            "request_body": {"projectName": "T727"},
            "response_body": '{"errno":0,"data":{"id":"estimate-1"}}',
        },
    )
    monkeypatch.setattr(
        client,
        "assign_testers",
        lambda estimate_id, tester_names: {
            "success": True,
            "status_code": 200,
                "request_body": {"id": estimate_id, "userIds": ["u1"]},
            "response_body": '{"errno":0}',
        },
    )

    result = client.run_create_then_assign({"testerNames": "TNNC01"})

    assert result["create_result"]["success"] is True
    assert result["assign_result"]["success"] is True
    assert result["assign_result"]["request_body"]["id"] == "estimate-1"


def test_run_create_then_assign_skips_assign_when_create_failed(monkeypatch):
    client = ManpowerApiClient(build_config(), SequenceAuthManager())
    monkeypatch.setattr(
        client,
        "create_estimate",
        lambda form_data: {"success": False, "status_code": 400, "request_body": None, "response_body": "bad"},
    )

    result = client.run_create_then_assign({"testerNames": "TNNC01"})

    assert result["create_result"]["success"] is False
    assert result["assign_result"]["skipped"] is True
    assert "未执行关联" in result["assign_result"]["response_body"]


def test_run_create_then_assign_skips_assign_when_tester_names_empty(monkeypatch):
    client = ManpowerApiClient(build_config(), SequenceAuthManager())
    monkeypatch.setattr(
        client,
        "create_estimate",
        lambda form_data: {
            "success": True,
            "status_code": 200,
            "request_body": {"projectName": "T727"},
            "response_body": '{"errno":0,"data":{"id":"estimate-1"}}',
        },
    )

    result = client.run_create_then_assign({"testerNames": " "})

    assert result["create_result"]["success"] is True
    assert result["assign_result"]["success"] is True
    assert result["assign_result"]["skipped"] is True
    assert result["assign_result"]["estimate_id"] == "estimate-1"


def test_run_create_then_assign_keeps_task_id_when_assign_failed(monkeypatch):
    client = ManpowerApiClient(build_config(), SequenceAuthManager())
    monkeypatch.setattr(
        client,
        "create_estimate",
        lambda form_data: {
            "success": True,
            "status_code": 200,
            "request_body": {"projectName": "T727"},
            "response_body": '{"errno":0,"data":{"id":"estimate-1"}}',
        },
    )
    monkeypatch.setattr(
        client,
        "assign_testers",
        lambda estimate_id, tester_names: {
            "success": False,
            "status_code": 200,
                "request_body": {"id": estimate_id, "userIds": ["u1"]},
            "response_body": "bad",
        },
    )

    result = client.run_create_then_assign({"testerNames": "TNNC01"})

    assert result["assign_result"]["success"] is False
    assert result["assign_result"]["estimate_id"] == "estimate-1"


def test_build_parser_supports_create_manpower():
    parser = workflow.build_parser()

    args = parser.parse_args(
        [
            "create-manpower",
            "--config",
            MANPOWER_CONFIG_PATH,
            "--form",
            MANPOWER_FORM_PATH,
        ]
    )

    assert args.command == "create-manpower"


def test_build_parser_supports_assign_and_run_manpower():
    parser = workflow.build_parser()

    assert parser.parse_args(["assign-manpower-testers"]).command == "assign-manpower-testers"
    assert parser.parse_args(["run-manpower"]).command == "run-manpower"


@pytest.mark.parametrize(
    ("command", "expected_config_path"),
    [
        ("create-manpower", workflow.DEFAULT_MANPOWER_CONFIG_PATH),
        ("discover-manpower", workflow.DEFAULT_MANPOWER_CONFIG_PATH),
        ("assign-manpower-testers", workflow.DEFAULT_MANPOWER_CONFIG_PATH),
        ("run-manpower", workflow.DEFAULT_MANPOWER_CONFIG_PATH),
        ("run", workflow.DEFAULT_CONFIG_PATH),
    ],
)
def test_resolve_config_path_uses_command_defaults(command, expected_config_path):
    parser = workflow.build_parser()
    args = parser.parse_args([command])

    assert Path(workflow._resolve_config_path(args)).resolve() == Path(expected_config_path).resolve()


def test_load_manpower_config_reads_example_file():
    config_path = Path(__file__).resolve().parents[1] / "examples" / "manpower_config.example.json"

    config = load_manpower_config(str(config_path))

    assert config.entry_url.endswith("#/manpowerhub/manpower/task?tab=taskhrmy")
    assert Path(config.capture.discovery_output_path).name == "manpower_task_capture.json"
    assert Path(config.capture.candidate_output_path).name == "manpower_task_candidate.json"


def test_main_dispatches_create_manpower(monkeypatch, capsys):
    calls = []

    class DummyAuthManager:
        def __init__(self, config):
            calls.append(("auth", config.base_url))

    class DummyClient:
        def __init__(self, config, auth_manager):
            calls.append(("client", config.base_url, isinstance(auth_manager, DummyAuthManager)))

        def create_estimate(self, form_data):
            calls.append(("create", form_data))
            return {"success": True, "status_code": 200, "request_body": form_data, "response_body": "ok"}

    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(
        workflow,
        "load_manpower_config",
        lambda path: SimpleNamespace(base_url="http://itms.tinno.com", config_path=path),
    )
    monkeypatch.setattr(workflow, "load_form_items", lambda args: [{"taskName": "接口提交"}])
    monkeypatch.setattr(workflow, "ItmsAuthManager", DummyAuthManager)
    monkeypatch.setattr(workflow, "ManpowerApiClient", DummyClient)
    monkeypatch.setattr(
        workflow.sys,
        "argv",
        [
            "main.py",
            "create-manpower",
            "--form",
            MANPOWER_FORM_PATH,
        ],
    )

    assert workflow.main() == 0
    assert calls == [
        ("auth", "http://itms.tinno.com"),
        ("client", "http://itms.tinno.com", True),
        ("create", {"taskName": "接口提交"}),
    ]
    assert '\"success\": true' in capsys.readouterr().out


def test_main_dispatches_assign_manpower_testers(monkeypatch, capsys):
    class DummyAuthManager:
        def __init__(self, config):
            self.config = config

    class DummyClient:
        def __init__(self, config, auth_manager):
            self.config = config
            self.auth_manager = auth_manager

        def find_estimate(self, form_data, estimate_id=None):
            return {"success": True, "estimate_id": "estimate-1", "matched_by": "estimateId"}

        def assign_testers(self, estimate_id, tester_names):
            return {"success": True, "request_body": {"id": estimate_id, "userIds": ["u1"]}, "response_body": "ok"}

    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(
        workflow,
        "load_manpower_config",
        lambda path: SimpleNamespace(base_url="http://itms.tinno.com", config_path=path),
    )
    monkeypatch.setattr(
        workflow,
        "load_form_items",
        lambda args: [{"projectName": "T727", "taskName": "X1103：专项验证", "beginDate": "2026-04-06", "endDate": "2026-04-12", "testerNames": "TNNC01"}],
    )
    monkeypatch.setattr(workflow, "ItmsAuthManager", DummyAuthManager)
    monkeypatch.setattr(workflow, "ManpowerApiClient", DummyClient)
    monkeypatch.setattr(workflow.sys, "argv", ["main.py", "assign-manpower-testers", "--form", MANPOWER_FORM_PATH])

    assert workflow.main() == 0
    output = capsys.readouterr().out
    assert '"estimate_id": "estimate-1"' in output
    assert '"userIds"' in output


def test_main_dispatches_run_manpower(monkeypatch, capsys):
    class DummyAuthManager:
        def __init__(self, config):
            self.config = config

    class DummyClient:
        def __init__(self, config, auth_manager):
            self.config = config
            self.auth_manager = auth_manager

        def run_create_then_assign(self, form_data):
            return {
                "create_result": {"success": True, "response_body": "create-ok"},
                "assign_result": {"success": True, "response_body": "assign-ok"},
            }

    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(
        workflow,
        "load_manpower_config",
        lambda path: SimpleNamespace(base_url="http://itms.tinno.com", config_path=path),
    )
    monkeypatch.setattr(workflow, "load_form_items", lambda args: [{"taskName": "一体执行", "testerNames": "TNNC01"}])
    monkeypatch.setattr(workflow, "ItmsAuthManager", DummyAuthManager)
    monkeypatch.setattr(workflow, "ManpowerApiClient", DummyClient)
    monkeypatch.setattr(workflow.sys, "argv", ["main.py", "run-manpower", "--form", MANPOWER_FORM_PATH])

    assert workflow.main() == 0
    output = capsys.readouterr().out
    assert '"create_result"' in output
    assert '"assign_result"' in output


def test_main_create_manpower_prints_single_result(monkeypatch, capsys):
    class DummyAuthManager:
        def __init__(self, config):
            self.config = config

    class DummyClient:
        def __init__(self, config, auth_manager):
            self.config = config
            self.auth_manager = auth_manager

        def create_estimate(self, form_data):
            return {"success": True, "taskName": form_data["taskName"], "status_code": 200}

    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(
        workflow,
        "load_manpower_config",
        lambda path: SimpleNamespace(base_url="http://itms.tinno.com", config_path=path),
    )
    monkeypatch.setattr(workflow, "load_form_items", lambda args: [{"taskName": "单条结果"}])
    monkeypatch.setattr(workflow, "ItmsAuthManager", DummyAuthManager)
    monkeypatch.setattr(workflow, "ManpowerApiClient", DummyClient)
    monkeypatch.setattr(
        workflow.sys,
        "argv",
        [
            "main.py",
            "create-manpower",
            "--form",
            MANPOWER_FORM_PATH,
        ],
    )

    assert workflow.main() == 0
    output = capsys.readouterr().out
    assert '"taskName": "单条结果"' in output
    assert output.strip().startswith("{")


def test_main_create_manpower_raises_when_form_items_missing(monkeypatch):
    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(
        workflow,
        "load_manpower_config",
        lambda path: SimpleNamespace(base_url="http://itms.tinno.com", config_path=path),
    )
    monkeypatch.setattr(workflow, "load_form_items", lambda args: [])
    monkeypatch.setattr(workflow, "ItmsAuthManager", lambda config: object())
    monkeypatch.setattr(workflow, "ManpowerApiClient", lambda config, auth_manager: object())
    monkeypatch.setattr(
        workflow.sys,
        "argv",
        ["main.py", "create-manpower"],
    )

    with pytest.raises(SystemExit, match="create-manpower"):
        workflow.main()


@pytest.mark.parametrize(
    ("command", "message"),
    [
        ("assign-manpower-testers", "assign-manpower-testers"),
        ("run-manpower", "run-manpower"),
    ],
)
def test_new_manpower_commands_require_form_items(monkeypatch, command, message):
    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(
        workflow,
        "load_manpower_config",
        lambda path: SimpleNamespace(base_url="http://itms.tinno.com", config_path=path),
    )
    monkeypatch.setattr(workflow, "load_form_items", lambda args: [])
    monkeypatch.setattr(workflow, "ItmsAuthManager", lambda config: object())
    monkeypatch.setattr(workflow, "ManpowerApiClient", lambda config, auth_manager: object())
    monkeypatch.setattr(workflow.sys, "argv", ["main.py", command])

    with pytest.raises(SystemExit, match=message):
        workflow.main()


@pytest.mark.parametrize("command", ["create-manpower", "discover-manpower"])
def test_manpower_commands_use_manpower_config_by_default(monkeypatch, command):
    calls = []

    class DummyAuthManager:
        def __init__(self, config):
            calls.append(("auth", config.config_path))

    class DummyDiscovery:
        def __init__(self, config, auth_manager):
            calls.append(("discovery", config.config_path, isinstance(auth_manager, DummyAuthManager)))

        def record_manual_flow(self):
            return {"candidate_request": {"command": command}}

    class DummyClient:
        def __init__(self, config, auth_manager):
            calls.append(("client", config.config_path, isinstance(auth_manager, DummyAuthManager)))

        def create_estimate(self, form_data):
            calls.append(("create", form_data))
            return {"success": True, "status_code": 200, "request_body": form_data, "response_body": "ok"}

    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(
        workflow,
        "load_manpower_config",
        lambda path: SimpleNamespace(base_url="http://itms.tinno.com", config_path=path),
    )
    monkeypatch.setattr(workflow, "load_form_items", lambda args: [{"taskName": "默认配置"}])
    monkeypatch.setattr(workflow, "ItmsAuthManager", DummyAuthManager)
    monkeypatch.setattr(workflow, "ManpowerCaptureDiscovery", DummyDiscovery)
    monkeypatch.setattr(workflow, "ManpowerApiClient", DummyClient)
    monkeypatch.setattr(
        workflow.sys,
        "argv",
        ["main.py", command, "--form", MANPOWER_FORM_PATH],
    )

    if command == "create-manpower":
        assert workflow.main() == 0
        assert any(item[0] == "client" and str(item[1]).replace("\\", "/").endswith(MANPOWER_CONFIG_PATH) and item[2] for item in calls)
        assert ("create", {"taskName": "默认配置"}) in calls
    else:
        assert workflow.main() == 0
        assert any(item[0] == "discovery" and str(item[1]).replace("\\", "/").endswith(MANPOWER_CONFIG_PATH) and item[2] for item in calls)


def test_main_create_manpower_returns_1_when_any_result_fails(monkeypatch, capsys):
    class DummyAuthManager:
        def __init__(self, config):
            self.config = config

    class DummyClient:
        def __init__(self, config, auth_manager):
            self.config = config
            self.auth_manager = auth_manager
            self.calls = 0

        def create_estimate(self, form_data):
            self.calls += 1
            if self.calls == 1:
                return {"success": True, "status_code": 200, "request_body": form_data, "response_body": "ok"}
            return {"success": False, "status_code": 400, "request_body": form_data, "response_body": "bad"}

    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(
        workflow,
        "load_manpower_config",
        lambda path: SimpleNamespace(base_url="http://itms.tinno.com", config_path=path),
    )
    monkeypatch.setattr(workflow, "load_form_items", lambda args: [{"taskName": "第一条"}, {"taskName": "第二条"}])
    monkeypatch.setattr(workflow, "ItmsAuthManager", DummyAuthManager)
    monkeypatch.setattr(workflow, "ManpowerApiClient", DummyClient)
    monkeypatch.setattr(
        workflow.sys,
        "argv",
        ["main.py", "create-manpower", "--form", MANPOWER_FORM_PATH],
    )

    assert workflow.main() == 1
    output = capsys.readouterr().out
    assert '"row_index": 1' in output
    assert '"row_index": 2' in output
    assert output.strip().startswith("[")


def test_main_create_manpower_returns_1_and_prints_standardized_result_when_client_raises(monkeypatch, capsys):
    class DummyAuthManager:
        def __init__(self, config):
            self.config = config

    class DummyClient:
        def __init__(self, config, auth_manager):
            self.config = config
            self.auth_manager = auth_manager

        def create_estimate(self, form_data):
            raise RuntimeError("network down")

        @staticmethod
        def build_failure_result(exc, request_body=None):
            return {
                "success": False,
                "status_code": None,
                "request_body": request_body,
                "response_body": str(exc),
            }

    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(
        workflow,
        "load_manpower_config",
        lambda path: SimpleNamespace(base_url="http://itms.tinno.com", config_path=path),
    )
    monkeypatch.setattr(workflow, "load_form_items", lambda args: [{"taskName": "异常提交"}])
    monkeypatch.setattr(workflow, "ItmsAuthManager", DummyAuthManager)
    monkeypatch.setattr(workflow, "ManpowerApiClient", DummyClient)
    monkeypatch.setattr(
        workflow.sys,
        "argv",
        ["main.py", "create-manpower", "--form", MANPOWER_FORM_PATH],
    )

    assert workflow.main() == 1
    output = capsys.readouterr().out
    assert '"success": false' in output
    assert '"status_code": null' in output
    assert '"request_body": null' in output
    assert '"response_body": "network down"' in output


def test_build_manpower_csv_from_weekly_plan_creates_expected_rows(tmp_path):
    source_path = tmp_path / "weekly_plan.xlsx"
    output_dir = tmp_path / "out"

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["项目", "本周计划", "人力投入", "项目编号", "测试人员"])
    sheet.append(["KO5", "第九轮monkey挂测，开关机，休眠唤醒，MTBF", 4, "V657", "张三,李四"])
    sheet.append(["lamu26", "MTBF测试收结果", 0.5, "P329D", "TNNC01"])
    workbook.save(source_path)

    output_path = build_manpower_csv_from_weekly_plan(
        str(source_path),
        output_dir=str(output_dir),
        today=date(2026, 4, 11),
        now=datetime(2026, 4, 11, 13, 45, 0),
    )

    rows = load_form_rows(str(output_path))
    assert rows == [
        {
            "projectName": "V657",
            "taskName": "KO5：第九轮monkey挂测，开关机，休眠唤醒，MTBF（WK15）",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026-04-06",
            "endDate": "2026-04-12",
            "standardWorkHours": "32",
            "testerNames": "张三,李四",
        },
        {
            "projectName": "P329D",
            "taskName": "lamu26：MTBF测试收结果（WK15）",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026-04-06",
            "endDate": "2026-04-12",
            "standardWorkHours": "4",
            "testerNames": "TNNC01",
        },
    ]
    assert Path(output_path).name == "manpower_form_template_20260411_134500.csv"


def test_build_manpower_csv_from_weekly_plan_uses_week_from_filename(tmp_path):
    source_path = tmp_path / "本周项目计划表_WK18.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["项目", "本周计划", "人力投入", "项目编号", "测试人员"])
    sheet.append(["KO5", "第十轮monkey专项挂测", 4, "V657", "张三"])
    workbook.save(source_path)

    output_path = build_manpower_csv_from_weekly_plan(
        str(source_path),
        output_dir=str(tmp_path),
        today=date(2026, 5, 6),
        now=datetime(2026, 5, 6, 15, 0, 0),
    )

    rows = load_form_rows(str(output_path))
    assert rows[0]["taskName"] == "KO5：第十轮monkey专项挂测（WK18）"
    assert rows[0]["beginDate"] == "2026-04-27"
    assert rows[0]["endDate"] == "2026-05-03"


def test_build_manpower_csv_from_weekly_plan_rejects_missing_headers(tmp_path):
    source_path = tmp_path / "weekly_plan.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["项目", "本周计划", "项目编号", "测试人员"])
    workbook.save(source_path)

    with pytest.raises(ValueError, match="缺少必要表头.*人力投入"):
        build_manpower_csv_from_weekly_plan(
            str(source_path),
            output_dir=str(tmp_path),
            today=date(2026, 4, 11),
            now=datetime(2026, 4, 11, 13, 45, 0),
        )


def test_build_manpower_csv_from_weekly_plan_rejects_non_excel_input(tmp_path):
    source_path = tmp_path / "weekly_plan.csv"
    source_path.write_text("项目,本周计划,人力投入,项目编号\n", encoding="utf-8")

    with pytest.raises(ValueError, match="只支持 Excel 文件"):
        build_manpower_csv_from_weekly_plan(
            str(source_path),
            output_dir=str(tmp_path),
            today=date(2026, 4, 11),
            now=datetime(2026, 4, 11, 13, 45, 0),
        )


def test_build_parser_supports_build_manpower_csv():
    parser = workflow.build_parser()

    args = parser.parse_args(["build-manpower-csv", "--excel", "examples/本周项目计划表.xlsx"])

    assert args.command == "build-manpower-csv"


def test_build_parser_supports_assign_manpower_from_weekly_plan():
    parser = workflow.build_parser()

    args = parser.parse_args(["assign-manpower-from-weekly-plan", "--excel", "examples/本周项目计划表.xlsx"])

    assert args.command == "assign-manpower-from-weekly-plan"


def test_main_dispatches_build_manpower_csv(monkeypatch, capsys):
    calls = []

    def fake_builder(source_path, output_dir):
        calls.append((source_path, output_dir))
        return "F:/automation-toolkit/python-tools/universal_Automation-Create-Task/examples/manpower_form_template_20260411_134500.csv"

    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(workflow, "build_manpower_csv_from_weekly_plan", fake_builder)
    monkeypatch.setattr(
        workflow.sys,
        "argv",
        ["main.py", "build-manpower-csv", "--excel", "examples/本周项目计划表.xlsx"],
    )

    assert workflow.main() == 0
    assert calls == [("examples/本周项目计划表.xlsx", str(workflow.TOOL_ROOT / "examples"))]
    assert "manpower_form_template_20260411_134500.csv" in capsys.readouterr().out


def test_main_dispatches_assign_manpower_from_weekly_plan(monkeypatch, capsys):
    calls = []

    class DummyAuthManager:
        def __init__(self, config):
            self.config = config
            calls.append(("auth", config.base_url))

    class DummyClient:
        def __init__(self, config, auth_manager):
            self.config = config
            self.auth_manager = auth_manager
            calls.append(("client", config.base_url, isinstance(auth_manager, DummyAuthManager)))

        def find_estimate(self, form_data, estimate_id=None):
            calls.append(("find", form_data, estimate_id))
            return {"success": True, "estimate_id": "estimate-1"}

        def assign_testers(self, estimate_id, tester_names):
            calls.append(("assign", estimate_id, tester_names))
            return {"success": True, "request_body": {"id": estimate_id, "userIds": ["u1"]}, "response_body": "ok"}

    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(
        workflow,
        "load_manpower_config",
        lambda path: SimpleNamespace(base_url="http://itms.tinno.com", config_path=path),
    )
    monkeypatch.setattr(workflow, "build_manpower_csv_from_weekly_plan", lambda source_path, output_dir: "generated.csv")
    monkeypatch.setattr(
        workflow,
        "load_form_rows",
        lambda path, sheet_name=None, header_row=1: [
            {"projectName": "T727", "taskName": "X1103：专项验证", "beginDate": "2026-04-06", "endDate": "2026-04-12", "testerNames": "TNNC01"}
        ],
    )
    monkeypatch.setattr(workflow, "ItmsAuthManager", DummyAuthManager)
    monkeypatch.setattr(workflow, "ManpowerApiClient", DummyClient)
    monkeypatch.setattr(
        workflow.sys,
        "argv",
        ["main.py", "assign-manpower-from-weekly-plan", "--excel", "examples/本周项目计划表.xlsx"],
    )

    assert workflow.main() == 0
    output = capsys.readouterr().out
    assert '"generated_csv_path": "generated.csv"' in output
    assert ("assign", "estimate-1", "TNNC01") in calls


def test_main_assign_manpower_from_weekly_plan_continues_when_find_estimate_fails(monkeypatch, capsys):
    class DummyAuthManager:
        def __init__(self, config):
            self.config = config

    class DummyClient:
        def __init__(self, config, auth_manager):
            self.calls = 0

        def find_estimate(self, form_data, estimate_id=None):
            self.calls += 1
            if self.calls == 1:
                return {"success": False, "error": "未找到匹配的人力预估任务"}
            return {"success": True, "estimate_id": "estimate-2"}

        def assign_testers(self, estimate_id, tester_names):
            return {"success": True, "request_body": {"id": estimate_id, "userIds": ["u2"]}, "response_body": "ok"}

    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(
        workflow,
        "load_manpower_config",
        lambda path: SimpleNamespace(base_url="http://itms.tinno.com", config_path=path),
    )
    monkeypatch.setattr(workflow, "build_manpower_csv_from_weekly_plan", lambda source_path, output_dir: "generated.csv")
    monkeypatch.setattr(
        workflow,
        "load_form_rows",
        lambda path, sheet_name=None, header_row=1: [
            {"projectName": "A", "taskName": "任务A", "beginDate": "2026-04-06", "endDate": "2026-04-12", "testerNames": "张三"},
            {"projectName": "B", "taskName": "任务B", "beginDate": "2026-04-06", "endDate": "2026-04-12", "testerNames": "李四"},
        ],
    )
    monkeypatch.setattr(workflow, "ItmsAuthManager", DummyAuthManager)
    monkeypatch.setattr(workflow, "ManpowerApiClient", DummyClient)
    monkeypatch.setattr(
        workflow.sys,
        "argv",
        ["main.py", "assign-manpower-from-weekly-plan", "--excel", "examples/本周项目计划表.xlsx"],
    )

    assert workflow.main() == 1
    output = capsys.readouterr().out
    assert '"row_index": 1' in output
    assert '"row_index": 2' in output


def test_build_parser_supports_fill_manpower_actual_from_weekly_plan():
    parser = workflow.build_parser()

    args = parser.parse_args(["fill-manpower-actual-from-weekly-plan", "--excel", "examples/本周项目计划表.xlsx"])

    assert args.command == "fill-manpower-actual-from-weekly-plan"


def test_build_parser_supports_complete_manpower():
    parser = workflow.build_parser()

    args = parser.parse_args(["complete-manpower", "--excel", "examples/manpower_form_template.csv"])

    assert args.command == "complete-manpower"


def test_build_parser_supports_complete_manpower_from_weekly_plan():
    parser = workflow.build_parser()

    args = parser.parse_args(["complete-manpower-from-weekly-plan", "--excel", "examples/本周项目计划表.xlsx"])

    assert args.command == "complete-manpower-from-weekly-plan"


def test_build_parser_supports_run_weekly_manpower():
    parser = workflow.build_parser()

    args = parser.parse_args(["run-weekly-manpower", "--excel", "examples/本周项目计划表.xlsx"])

    assert args.command == "run-weekly-manpower"


def test_build_parser_supports_monday_manpower_maintenance():
    parser = workflow.build_parser()

    args = parser.parse_args(["monday-manpower-maintenance", "--excel", "examples/本周项目计划表.xlsx"])

    assert args.command == "monday-manpower-maintenance"


def test_main_dispatches_fill_manpower_actual_from_weekly_plan(monkeypatch, capsys):
    calls = []

    class DummyAuthManager:
        def __init__(self, config):
            self.config = config
            calls.append(("auth", config.base_url))

    class DummyClient:
        def __init__(self, config, auth_manager):
            calls.append(("client", config.base_url, isinstance(auth_manager, DummyAuthManager)))

        def find_estimate(self, form_data, estimate_id=None):
            calls.append(("find", form_data, estimate_id))
            return {"success": True, "estimate_id": "estimate-1"}

        def fill_actual_work(self, estimate_id, effort, tester_names, actual_ratios):
            calls.append(("fill", estimate_id, effort, tester_names, actual_ratios))
            return {"success": True, "request_body": {"id": estimate_id}, "response_body": "ok"}

    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(
        workflow,
        "load_manpower_config",
        lambda path: SimpleNamespace(base_url="http://itms.tinno.com", config_path=path),
    )
    monkeypatch.setattr(
        workflow,
        "load_weekly_plan_actual_rows",
        lambda source_path: [
            {
                "projectName": "V551A",
                "taskName": "V551A：专项验证",
                "beginDate": "2026-04-06",
                "endDate": "2026-04-12",
                "effort": 8,
                "testerNames": ["吕代"],
                "actualRatios": [0.1],
            }
        ],
    )
    monkeypatch.setattr(workflow, "ItmsAuthManager", DummyAuthManager)
    monkeypatch.setattr(workflow, "ManpowerApiClient", DummyClient)
    monkeypatch.setattr(
        workflow.sys,
        "argv",
        ["main.py", "fill-manpower-actual-from-weekly-plan", "--excel", "examples/本周项目计划表.xlsx"],
    )

    assert workflow.main() == 0
    output = capsys.readouterr().out
    assert '"row_index": 1' in output
    assert ("fill", "estimate-1", 8, ["吕代"], [0.1]) in calls


def test_main_dispatches_complete_manpower(monkeypatch, capsys):
    class DummyAuthManager:
        def __init__(self, config):
            self.config = config

    class DummyClient:
        def __init__(self, config, auth_manager):
            pass

        def find_estimate(self, form_data, estimate_id=None):
            return {"success": True, "estimate_id": "estimate-1"}

        def complete_estimate(self, estimate_id):
            return {"success": True, "request_body": {"ids": estimate_id, "status": 2}, "response_body": "ok"}

    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(workflow, "load_manpower_config", lambda path: SimpleNamespace(base_url="http://itms.tinno.com"))
    monkeypatch.setattr(
        workflow,
        "load_form_rows",
        lambda path, sheet_name=None, header_row=1: [{"projectName": "平台类", "taskName": "任务A", "beginDate": "2026-04-06", "endDate": "2026-04-12"}],
    )
    monkeypatch.setattr(workflow, "ItmsAuthManager", DummyAuthManager)
    monkeypatch.setattr(workflow, "ManpowerApiClient", DummyClient)
    monkeypatch.setattr(workflow.sys, "argv", ["main.py", "complete-manpower", "--excel", "examples/manpower_form_template.csv"])

    assert workflow.main() == 0
    assert '"status": 2' in capsys.readouterr().out


def test_main_dispatches_complete_manpower_from_weekly_plan(monkeypatch, capsys):
    class DummyAuthManager:
        def __init__(self, config):
            self.config = config

    class DummyClient:
        def __init__(self, config, auth_manager):
            pass

        def find_estimate(self, form_data, estimate_id=None):
            return {"success": True, "estimate_id": "estimate-9"}

        def complete_estimate(self, estimate_id):
            return {"success": True, "request_body": {"ids": estimate_id, "status": 2}, "response_body": "ok"}

    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(workflow, "load_manpower_config", lambda path: SimpleNamespace(base_url="http://itms.tinno.com"))
    monkeypatch.setattr(workflow, "build_manpower_csv_from_weekly_plan", lambda source_path, output_dir: "generated.csv")
    monkeypatch.setattr(
        workflow,
        "load_form_rows",
        lambda path, sheet_name=None, header_row=1: [{"projectName": "平台类", "taskName": "任务A", "beginDate": "2026-04-06", "endDate": "2026-04-12"}],
    )
    monkeypatch.setattr(workflow, "ItmsAuthManager", DummyAuthManager)
    monkeypatch.setattr(workflow, "ManpowerApiClient", DummyClient)
    monkeypatch.setattr(
        workflow.sys,
        "argv",
        ["main.py", "complete-manpower-from-weekly-plan", "--excel", "examples/本周项目计划表.xlsx"],
    )

    assert workflow.main() == 0
    output = capsys.readouterr().out
    assert '"generated_csv_path": "generated.csv"' in output
    assert '"status": 2' in output


def test_main_dispatches_run_weekly_manpower(monkeypatch, capsys):
    class DummyAuthManager:
        def __init__(self, config):
            self.config = config

    class DummyClient:
        def __init__(self, config, auth_manager):
            pass

        def create_estimate(self, form_data):
            return {"success": True, "response_body": '{"errno":0,"data":{"id":"estimate-1"}}'}

        def assign_testers(self, estimate_id, tester_names):
            return {"success": True, "response_body": "assign-ok"}

        def fill_actual_work(self, estimate_id, effort, tester_names, actual_ratios):
            return {"success": True, "response_body": "actual-ok"}

    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(workflow, "load_manpower_config", lambda path: SimpleNamespace(base_url="http://itms.tinno.com"))
    monkeypatch.setattr(workflow, "build_manpower_csv_from_weekly_plan", lambda source_path, output_dir: "generated.csv")
    monkeypatch.setattr(
        workflow,
        "load_form_rows",
        lambda path, sheet_name=None, header_row=1: [{"projectName": "平台类", "taskName": "任务A", "beginDate": "2026-04-06", "endDate": "2026-04-12", "testerNames": "洪子鑫"}],
    )
    monkeypatch.setattr(
        workflow,
        "load_weekly_plan_actual_rows",
        lambda source_path: [{"projectName": "平台类", "taskName": "任务A", "beginDate": "2026-04-06", "endDate": "2026-04-12", "effort": 6, "testerNames": ["洪子鑫"], "actualRatios": [1]}],
    )
    monkeypatch.setattr(workflow, "ItmsAuthManager", DummyAuthManager)
    monkeypatch.setattr(workflow, "ManpowerApiClient", DummyClient)
    monkeypatch.setattr(workflow.sys, "argv", ["main.py", "run-weekly-manpower", "--excel", "examples/本周项目计划表.xlsx"])

    assert workflow.main() == 0
    output = capsys.readouterr().out
    assert '"create_result"' in output
    assert '"assign_result"' in output
    assert '"fill_actual_result"' in output


def test_main_run_weekly_manpower_skips_later_steps_when_create_failed(monkeypatch, capsys):
    class DummyAuthManager:
        def __init__(self, config):
            self.config = config

    class DummyClient:
        def __init__(self, config, auth_manager):
            pass

        def create_estimate(self, form_data):
            return {"success": False, "response_body": "bad"}

    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(workflow, "load_manpower_config", lambda path: SimpleNamespace(base_url="http://itms.tinno.com"))
    monkeypatch.setattr(workflow, "build_manpower_csv_from_weekly_plan", lambda source_path, output_dir: "generated.csv")
    monkeypatch.setattr(
        workflow,
        "load_form_rows",
        lambda path, sheet_name=None, header_row=1: [{"projectName": "平台类", "taskName": "任务A", "beginDate": "2026-04-06", "endDate": "2026-04-12", "testerNames": "洪子鑫"}],
    )
    monkeypatch.setattr(
        workflow,
        "load_weekly_plan_actual_rows",
        lambda source_path: [{"projectName": "平台类", "taskName": "任务A", "beginDate": "2026-04-06", "endDate": "2026-04-12", "effort": 6, "testerNames": ["洪子鑫"], "actualRatios": [1]}],
    )
    monkeypatch.setattr(workflow, "ItmsAuthManager", DummyAuthManager)
    monkeypatch.setattr(workflow, "ManpowerApiClient", DummyClient)
    monkeypatch.setattr(workflow.sys, "argv", ["main.py", "run-weekly-manpower", "--excel", "examples/本周项目计划表.xlsx"])

    assert workflow.main() == 1
    output = capsys.readouterr().out
    assert "创建失败，未执行关联" in output
    assert "创建失败，未执行实际工时回填" in output


def test_main_dispatches_monday_manpower_maintenance(monkeypatch, capsys):
    class DummyAuthManager:
        def __init__(self, config):
            self.config = config

    class DummyClient:
        def __init__(self, config, auth_manager):
            pass

        def find_estimate(self, form_data, estimate_id=None):
            return {"success": True, "estimate_id": "estimate-1"}

        def complete_estimate(self, estimate_id):
            return {"success": True, "request_body": {"ids": estimate_id, "status": 2}, "response_body": "ok"}

    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(workflow, "load_manpower_config", lambda path: SimpleNamespace(base_url="http://itms.tinno.com"))
    monkeypatch.setattr(workflow, "build_manpower_csv_from_weekly_plan", lambda source_path, output_dir: "generated.csv")
    monkeypatch.setattr(
        workflow,
        "load_form_rows",
        lambda path, sheet_name=None, header_row=1: [{"projectName": "平台类", "taskName": "任务A", "beginDate": "2026-04-06", "endDate": "2026-04-12"}],
    )
    monkeypatch.setattr(workflow, "ItmsAuthManager", DummyAuthManager)
    monkeypatch.setattr(workflow, "ManpowerApiClient", DummyClient)
    monkeypatch.setattr(workflow.sys, "argv", ["main.py", "monday-manpower-maintenance", "--excel", "examples/本周项目计划表.xlsx"])

    assert workflow.main() == 0
    output = capsys.readouterr().out
    assert '"reminder_message"' in output
    assert "请编写并确认本周项目计划表.xlsx" in output


def test_main_build_manpower_csv_requires_excel(monkeypatch):
    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(workflow.sys, "argv", ["main.py", "build-manpower-csv"])

    with pytest.raises(SystemExit, match="build-manpower-csv"):
        workflow.main()


def test_main_assign_manpower_from_weekly_plan_requires_excel(monkeypatch):
    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(workflow.sys, "argv", ["main.py", "assign-manpower-from-weekly-plan"])

    with pytest.raises(SystemExit, match="assign-manpower-from-weekly-plan"):
        workflow.main()


def test_main_fill_manpower_actual_from_weekly_plan_requires_excel(monkeypatch):
    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(workflow.sys, "argv", ["main.py", "fill-manpower-actual-from-weekly-plan"])

    with pytest.raises(SystemExit, match="fill-manpower-actual-from-weekly-plan"):
        workflow.main()


def test_main_complete_manpower_requires_excel_or_form(monkeypatch):
    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(workflow.sys, "argv", ["main.py", "complete-manpower"])

    with pytest.raises(SystemExit, match="complete-manpower"):
        workflow.main()


def test_main_complete_manpower_from_weekly_plan_requires_excel(monkeypatch):
    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(workflow.sys, "argv", ["main.py", "complete-manpower-from-weekly-plan"])

    with pytest.raises(SystemExit, match="complete-manpower-from-weekly-plan"):
        workflow.main()


def test_main_run_weekly_manpower_requires_excel(monkeypatch):
    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(workflow.sys, "argv", ["main.py", "run-weekly-manpower"])

    with pytest.raises(SystemExit, match="run-weekly-manpower"):
        workflow.main()


def test_main_monday_manpower_maintenance_requires_excel(monkeypatch):
    monkeypatch.setattr(workflow, "setup_logging", lambda: None)
    monkeypatch.setattr(workflow.sys, "argv", ["main.py", "monday-manpower-maintenance"])

    with pytest.raises(SystemExit, match="monday-manpower-maintenance"):
        workflow.main()


def test_readme_mentions_build_manpower_csv_command():
    readme_path = Path(__file__).resolve().parents[1] / "README.md"
    content = readme_path.read_text(encoding="utf-8")

    assert "build-manpower-csv" in content
    assert "本周项目计划表.xlsx" in content
    assert "manpower_form_template_" in content


def test_readme_mentions_assign_and_run_manpower_commands():
    readme_path = Path(__file__).resolve().parents[1] / "README.md"
    content = readme_path.read_text(encoding="utf-8")

    assert "assign-manpower-testers" in content
    assert "run-manpower" in content
    assert "testerNames" in content


def test_readme_mentions_assign_manpower_from_weekly_plan_command():
    readme_path = Path(__file__).resolve().parents[1] / "README.md"
    content = readme_path.read_text(encoding="utf-8")

    assert "assign-manpower-from-weekly-plan" in content
    assert "本周项目计划表.xlsx" in content


def test_readme_mentions_fill_manpower_actual_from_weekly_plan_command():
    readme_path = Path(__file__).resolve().parents[1] / "README.md"
    content = readme_path.read_text(encoding="utf-8")

    assert "fill-manpower-actual-from-weekly-plan" in content
    assert "实际工时01" in content


def test_readme_mentions_complete_manpower_commands():
    readme_path = Path(__file__).resolve().parents[1] / "README.md"
    content = readme_path.read_text(encoding="utf-8")

    assert "complete-manpower" in content
    assert "complete-manpower-from-weekly-plan" in content
    assert "status=2" in content


def test_readme_mentions_weekly_orchestration_commands():
    readme_path = Path(__file__).resolve().parents[1] / "README.md"
    content = readme_path.read_text(encoding="utf-8")

    assert "run-weekly-manpower" in content
    assert "monday-manpower-maintenance" in content
    assert "请编写并确认本周项目计划表.xlsx" in content


def test_build_manpower_csv_from_weekly_plan_writes_tester_names(tmp_path):
    source_path = tmp_path / "weekly_plan.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["项目", "本周计划", "人力投入", "项目编号", "测试人员"])
    sheet.append(["X1103", "专项验证", 2, "T727", "张三,李四"])
    workbook.save(source_path)

    output_path = build_manpower_csv_from_weekly_plan(
        str(source_path),
        output_dir=str(tmp_path),
        today=date(2026, 4, 11),
        now=datetime(2026, 4, 11, 16, 0, 0),
    )

    with open(output_path, "r", encoding="utf-8-sig", newline="") as file_obj:
        rows = list(__import__("csv").DictReader(file_obj))

    assert rows[0]["testerNames"] == "张三,李四"


def test_build_manpower_csv_from_weekly_plan_ignores_incomplete_rows(tmp_path):
    source_path = tmp_path / "weekly_plan.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["项目", "本周计划", "人力投入", "项目编号", "测试人员"])
    sheet.append(["KO5", "专项验证", 4, "V657", "张三"])
    sheet.append([None, "残留说明文本", None, None, None])
    workbook.save(source_path)

    output_path = build_manpower_csv_from_weekly_plan(
        str(source_path),
        output_dir=str(tmp_path),
        today=date(2026, 4, 11),
        now=datetime(2026, 4, 11, 16, 0, 0),
    )

    rows = load_form_rows(str(output_path))
    assert rows == [
        {
            "projectName": "V657",
            "taskName": "KO5：专项验证（WK15）",
            "site": "南昌",
            "groupType": "系统",
            "beginDate": "2026-04-06",
            "endDate": "2026-04-12",
            "standardWorkHours": "32",
            "testerNames": "张三",
        }
    ]


def test_load_weekly_plan_actual_rows_reads_testers_and_actual_columns(tmp_path):
    source_path = tmp_path / "weekly_plan.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["项目", "本周计划", "人力投入", "项目编号", "测试人员", "实际工时01", "实际工时02"])
    sheet.append(["V551A", "专项验证", 8, "V551A", "吕代,张三", 0.1, 0.2])
    workbook.save(source_path)

    rows = load_weekly_plan_actual_rows(str(source_path), today=date(2026, 4, 11))

    assert rows == [
        {
            "projectName": "V551A",
            "taskName": "V551A：专项验证（WK15）",
            "beginDate": "2026-04-06",
            "endDate": "2026-04-12",
            "effort": 8,
            "testerNames": ["吕代", "张三"],
            "actualRatios": [0.1, 0.2],
        }
    ]


def test_load_weekly_plan_actual_rows_uses_week_from_filename(tmp_path):
    source_path = tmp_path / "本周项目计划表_WK18.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["项目", "本周计划", "人力投入", "项目编号", "测试人员", "实际工时01"])
    sheet.append(["KO5", "第十轮monkey专项挂测", 4, "V657", "张三", 1])
    workbook.save(source_path)

    rows = load_weekly_plan_actual_rows(str(source_path), today=date(2026, 5, 6))

    assert rows[0]["taskName"] == "KO5：第十轮monkey专项挂测（WK18）"
    assert rows[0]["beginDate"] == "2026-04-27"
    assert rows[0]["endDate"] == "2026-05-03"


def test_load_weekly_plan_actual_rows_ignores_incomplete_rows(tmp_path):
    source_path = tmp_path / "weekly_plan.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["项目", "本周计划", "人力投入", "项目编号", "测试人员", "实际工时01"])
    sheet.append(["V551A", "专项验证", 8, "V551A", "吕代", 0.1])
    sheet.append([None, "残留说明文本", None, None, None, None])
    workbook.save(source_path)

    rows = load_weekly_plan_actual_rows(str(source_path), today=date(2026, 4, 11))

    assert rows == [
        {
            "projectName": "V551A",
            "taskName": "V551A：专项验证（WK15）",
            "beginDate": "2026-04-06",
            "endDate": "2026-04-12",
            "effort": 8,
            "testerNames": ["吕代"],
            "actualRatios": [0.1],
        }
    ]
