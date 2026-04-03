from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .models import ProjectInfo, ReportInput, SpecialtyExecution


REQUIRED_SHEETS = ("项目基础信息", "专项执行清单")
PROJECT_REQUIRED_COLUMNS = (
    "project_name",
    "project_stage",
    "software_version",
    "reporter",
    "report_date",
    "output_root",
    "template_profile",
)
EXECUTION_REQUIRED_COLUMNS = (
    "specialty_code",
    "specialty_name",
    "enabled",
    "round_no",
    "test_cycle",
    "test_result",
    "device_count",
    "bug_summary",
    "subreport_title",
)


class InputValidationError(ValueError):
    """输入工作簿校验异常。"""


def _normalize_header(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _sheet_rows(worksheet) -> list[dict[str, object]]:
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [_normalize_header(cell) for cell in rows[0]]
    parsed_rows: list[dict[str, object]] = []
    for raw_row in rows[1:]:
        if raw_row is None or all(cell in (None, "") for cell in raw_row):
            continue
        parsed_rows.append({
            headers[index]: raw_row[index] if index < len(raw_row) else None
            for index in range(len(headers))
            if headers[index]
        })
    return parsed_rows


def _require_columns(rows: list[dict[str, object]], required_columns: tuple[str, ...], sheet_name: str) -> None:
    if not rows:
        raise InputValidationError(f"{sheet_name} 为空，无法继续生成报告")

    row_keys = set(rows[0].keys())
    missing = [column for column in required_columns if column not in row_keys]
    if missing:
        raise InputValidationError(f"{sheet_name} 缺少必要列: {', '.join(missing)}")


def _to_bool(value: object) -> bool:
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    return normalized in {"1", "true", "yes", "y", "是"}


def _to_str(value: object) -> str:
    return "" if value is None else str(value).strip()


def _to_int(value: object, field_name: str, specialty_code: str) -> int:
    try:
        return int(value)
    except (TypeError, ValueError) as exc:
        raise InputValidationError(f"专项 {specialty_code} 的字段 {field_name} 必须是整数") from exc


def read_input_workbook(workbook_path: Path | str) -> ReportInput:
    workbook_path = Path(workbook_path)
    if not workbook_path.exists():
        raise FileNotFoundError(f"输入工作簿不存在: {workbook_path}")

    workbook = load_workbook(workbook_path, data_only=True)

    missing_sheets = [sheet_name for sheet_name in REQUIRED_SHEETS if sheet_name not in workbook.sheetnames]
    if missing_sheets:
        raise InputValidationError(f"输入工作簿缺少必要工作表: {', '.join(missing_sheets)}")

    project_rows = _sheet_rows(workbook["项目基础信息"])
    _require_columns(project_rows, PROJECT_REQUIRED_COLUMNS, "项目基础信息")
    project_row = project_rows[0]
    project = ProjectInfo(
        project_name=_to_str(project_row["project_name"]),
        project_stage=_to_str(project_row["project_stage"]),
        software_version=_to_str(project_row["software_version"]),
        reporter=_to_str(project_row["reporter"]),
        report_date=_to_str(project_row["report_date"]),
        output_root=Path(_to_str(project_row["output_root"])),
        template_profile=_to_str(project_row["template_profile"]),
        brand=_to_str(project_row.get("brand")),
        platform=_to_str(project_row.get("platform")),
        bom=_to_str(project_row.get("bom")),
        rom_ram=_to_str(project_row.get("rom_ram")),
    )

    execution_rows = _sheet_rows(workbook["专项执行清单"])
    _require_columns(execution_rows, EXECUTION_REQUIRED_COLUMNS, "专项执行清单")

    extension_by_specialty: dict[str, dict[str, str]] = {}
    if "专项扩展字段" in workbook.sheetnames:
        extension_rows = _sheet_rows(workbook["专项扩展字段"])
        for row in extension_rows:
            specialty_code = _to_str(row.get("specialty_code"))
            field_key = _to_str(row.get("field_key"))
            field_value = _to_str(row.get("field_value"))
            if not specialty_code or not field_key:
                continue
            extension_by_specialty.setdefault(specialty_code, {})[field_key] = field_value

    specialties: list[SpecialtyExecution] = []
    for row in execution_rows:
        specialty_code = _to_str(row["specialty_code"])
        specialties.append(
            SpecialtyExecution(
                specialty_code=specialty_code,
                specialty_name=_to_str(row["specialty_name"]),
                enabled=_to_bool(row["enabled"]),
                round_no=_to_int(row["round_no"], "round_no", specialty_code),
                test_cycle=_to_str(row["test_cycle"]),
                test_result=_to_str(row["test_result"]),
                device_count=_to_str(row["device_count"]),
                bug_summary=_to_str(row["bug_summary"]),
                subreport_title=_to_str(row["subreport_title"]),
                extensions=extension_by_specialty.get(specialty_code, {}),
            )
        )

    bug_rows: list[dict[str, str]] = []
    if "Bug汇总导入" in workbook.sheetnames:
        for row in _sheet_rows(workbook["Bug汇总导入"]):
            bug_rows.append({key: _to_str(value) for key, value in row.items()})

    return ReportInput(project=project, specialties=specialties, bug_rows=bug_rows)
