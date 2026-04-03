from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from .config import get_specialty_configs
from .models import TemplatePaths


SUMMARY_HEADERS = [
    "Issue key",
    "Project",
    "Affects Version/s",
    "Component/s",
    "Priority",
    "Custom field (Risk)",
    "Summary",
    "Description",
    "Reporter",
]


def _populate_overview_sheet(worksheet) -> None:
    worksheet.title = "总览"
    worksheet["A1"] = "稳定性测试报告汇总"
    worksheet["A2"] = "项目"
    worksheet["C2"] = "阶段"
    worksheet["A3"] = "版本"
    worksheet["C3"] = "报告人"
    worksheet["A5"] = "专项"
    worksheet["B5"] = "测试结果"
    worksheet["C5"] = "轮次"
    worksheet["D5"] = "报告文件"


def _populate_specialty_sheet(worksheet, display_name: str) -> None:
    worksheet["A1"] = f"{display_name} 专项测试报告"
    worksheet["A2"] = "测试类型"
    worksheet["C2"] = "当前轮次"
    worksheet["A3"] = "测试阶段"
    worksheet["C3"] = "测试结果"
    worksheet["A4"] = "报告人"
    worksheet["C4"] = "测试时间"
    worksheet["A5"] = "测试版本"
    worksheet["C5"] = "测试样机数量"
    worksheet["A6"] = "测试概览"
    worksheet["A8"] = "测试轮数"
    worksheet["B8"] = "测试日期"
    worksheet["C8"] = "测试版本"
    worksheet["D8"] = "Result"


def _populate_bug_sheet(worksheet) -> None:
    worksheet.title = "Bug_MonkeyList"
    worksheet.append(SUMMARY_HEADERS)


def _populate_appendix_sheet(worksheet) -> None:
    worksheet.title = "附录-问题等级标准"
    worksheet.append(["附录：问题等级标准", "", "", "", ""])
    worksheet.append(["问题等级", "报错类型", "次数要求", "报错解释", ""])
    worksheet.append(["A", "SWT", "无次数要求", "System Server Watchdog，对应现象重启", ""])
    worksheet.append(["A", "fatal.JE", "无次数要求", "System Server Java 异常，对应现象重启", ""])
    worksheet.append(["B", "ANR", "按专项标准统计", "应用无响应", ""])
    worksheet.append(["C", "Crash", "按专项标准统计", "应用崩溃", ""])


def _create_summary_template(summary_path: Path) -> None:
    workbook = Workbook()
    _populate_overview_sheet(workbook.active)
    for specialty in get_specialty_configs().values():
        specialty_sheet = workbook.create_sheet(specialty.summary_sheet_name)
        _populate_specialty_sheet(specialty_sheet, specialty.display_name)
    bug_sheet = workbook.create_sheet("Bug_MonkeyList")
    _populate_bug_sheet(bug_sheet)
    appendix_sheet = workbook.create_sheet("附录-问题等级标准")
    _populate_appendix_sheet(appendix_sheet)
    workbook.save(summary_path)


def _create_specialty_template(template_path: Path, display_name: str) -> None:
    workbook = Workbook()
    report_sheet = workbook.active
    report_sheet.title = "报告"
    _populate_specialty_sheet(report_sheet, display_name)
    workbook.save(template_path)


def ensure_builtin_templates(templates_root: Path) -> TemplatePaths:
    templates_root.mkdir(parents=True, exist_ok=True)

    summary_path = templates_root / "stability/summary/stability_summary.xlsx"
    specialty_root = templates_root / "stability/specialties"
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    specialty_root.mkdir(parents=True, exist_ok=True)

    if not summary_path.exists():
        _create_summary_template(summary_path)

    created_templates: set[Path] = set()
    for specialty in get_specialty_configs().values():
        template_path = templates_root / specialty.template_relative_path
        template_path.parent.mkdir(parents=True, exist_ok=True)
        if template_path.exists() or template_path in created_templates:
            continue
        _create_specialty_template(template_path, specialty.display_name)
        created_templates.add(template_path)

    return TemplatePaths(summary=summary_path)


def create_input_template(output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()

    project_sheet = workbook.active
    project_sheet.title = "项目基础信息"
    project_sheet.append([
        "project_name",
        "project_stage",
        "brand",
        "platform",
        "bom",
        "rom_ram",
        "software_version",
        "reporter",
        "report_date",
        "output_root",
        "template_profile",
    ])
    project_sheet.append([
        "KO5",
        "Beta",
        "TECNO",
        "MT6789",
        "P1",
        "256G+8G",
        "KO5-16.2.0.116(OP001PF001AZ)_SU",
        "李小红",
        "2026-03-17",
        str(output_path.parent / "output"),
        "default",
    ])

    execution_sheet = workbook.create_sheet("专项执行清单")
    execution_sheet.append([
        "specialty_code",
        "specialty_name",
        "enabled",
        "round_no",
        "test_cycle",
        "test_result",
        "device_count",
        "bug_summary",
        "subreport_title",
    ])
    execution_sheet.append([
        "mtbf",
        "MTBF",
        "TRUE",
        "5",
        "2026/3/10-2026/3/17",
        "PASS",
        "10",
        "Blocker：0、Critical：0、Major：0",
        "KO5项目Beta阶段MTBF专项第五轮测试报告",
    ])
    execution_sheet.append([
        "stress",
        "Stress",
        "TRUE",
        "3",
        "2026/3/10-2026/3/17",
        "PASS",
        "20",
        "Blocker：0、Critical：0、Major：0",
        "KO5项目Beta阶段Stress专项第三轮测试报告",
    ])

    extension_sheet = workbook.create_sheet("专项扩展字段")
    extension_sheet.append(["specialty_code", "field_key", "field_value"])
    extension_sheet.append(["stress", "stress_summary", "KO5压力稳定性第3轮测试PASS，此轮发现问题0个"])
    extension_sheet.append(["stress", "sample_count", "20"])

    bug_sheet = workbook.create_sheet("Bug汇总导入")
    bug_sheet.append(SUMMARY_HEADERS)
    bug_sheet.append([
        "KO5OS16AEE-748",
        "KO5",
        "KO5-16.2.0.116(OP001PF001AZ)_SU",
        "Reboot",
        "Blocker",
        "occasional",
        "[MonkeyAEE] SWT system_server",
        "示例问题",
        "xiaohongli.tinno",
    ])

    workbook.save(output_path)
    return output_path
