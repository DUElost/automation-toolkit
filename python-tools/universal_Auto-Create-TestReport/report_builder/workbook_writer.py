from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .models import ProjectInfo, SpecialtyConfig, SpecialtyExecution
from .title_helpers import build_exact_report_title, build_summary_title, chinese_round_text, numeric_round_text


def _round_label(round_no: int) -> str:
    return f"第{round_no}轮"


def _extension_summary(specialty: SpecialtyExecution) -> str:
    return specialty.extensions.get("stress_summary") or specialty.extensions.get("summary") or specialty.bug_summary


def _write_common_fields(worksheet, project: ProjectInfo, specialty: SpecialtyExecution, config: SpecialtyConfig) -> None:
    worksheet["A1"] = specialty.subreport_title or f"{project.project_name} {config.display_name} 专项测试报告"
    worksheet["B2"] = config.display_name
    worksheet["D2"] = _round_label(specialty.round_no)
    worksheet["B3"] = project.project_stage
    worksheet["D3"] = specialty.test_result
    worksheet["B4"] = project.reporter
    worksheet["D4"] = specialty.test_cycle
    worksheet["B5"] = project.software_version
    worksheet["D5"] = specialty.extensions.get("sample_count") or specialty.device_count
    worksheet["B6"] = _extension_summary(specialty)

    history_row = 9
    while worksheet[f"A{history_row}"].value not in (None, ""):
        history_row += 1

    worksheet[f"A{history_row}"] = _round_label(specialty.round_no)
    worksheet[f"B{history_row}"] = specialty.test_cycle
    worksheet[f"C{history_row}"] = project.software_version
    worksheet[f"D{history_row}"] = specialty.test_result

    extension_row = 12
    if specialty.extensions:
        worksheet[f"A{extension_row}"] = "扩展字段"
        for key, value in specialty.extensions.items():
            extension_row += 1
            worksheet[f"A{extension_row}"] = key
            worksheet[f"B{extension_row}"] = value


def write_specialty_report(report_path: Path, project: ProjectInfo, specialty: SpecialtyExecution, config: SpecialtyConfig) -> None:
    workbook = load_workbook(report_path)
    exact_title = build_exact_report_title(project, specialty, config)
    if "报告" in workbook.sheetnames:
        worksheet = workbook["报告"]
        _write_common_fields(worksheet, project, specialty, config)
    elif "测试报告" in workbook.sheetnames:
        worksheet = workbook["测试报告"]
        existing_title = worksheet["B2"].value or ""
        existing_title_lines = str(existing_title).split("\n")
        if len(existing_title_lines) > 1:
            worksheet["B2"] = f"{exact_title}\n{existing_title_lines[1]}"
        else:
            worksheet["B2"] = exact_title
        worksheet["C3"] = specialty.test_result
        worksheet["C4"] = project.project_name
        worksheet["H4"] = project.project_stage
        worksheet["C5"] = project.software_version
        worksheet["E5"] = specialty.test_cycle
        worksheet["H5"] = project.reporter
        try:
            worksheet["C6"] = int(str(specialty.extensions.get("sample_count") or specialty.device_count))
        except ValueError:
            worksheet["C6"] = specialty.extensions.get("sample_count") or specialty.device_count
        worksheet["G7"] = 0 if specialty.test_result.upper() == "PASS" else worksheet["G7"].value
        worksheet["H7"] = specialty.bug_summary.replace("、", "\n").replace("：", "：")
        worksheet["B12"] = specialty.extensions.get("stress_summary") or (
            f"{exact_title}：\n本轮测试结果为{specialty.test_result}。"
        )
    elif "result" in workbook.sheetnames:
        worksheet = workbook["result"]
        worksheet["A1"] = exact_title
        worksheet["D4"] = project.project_name
        worksheet["F4"] = project.project_stage
        worksheet["B6"] = project.software_version
        worksheet["H6"] = project.reporter
        worksheet["F5"] = worksheet["F5"].value or worksheet["F5"].value
    workbook.save(report_path)


def _write_overview_sheet(worksheet, project: ProjectInfo, specialties: list[tuple[SpecialtyExecution, SpecialtyConfig, Path]]) -> None:
    worksheet["A1"] = f"{project.project_name}项目{project.project_stage}阶段稳定性测试报告汇总"
    worksheet["B2"] = project.project_name
    worksheet["D2"] = project.project_stage
    worksheet["B3"] = project.software_version
    worksheet["D3"] = project.reporter

    row_index = 6
    for specialty, config, report_path in specialties:
        worksheet[f"A{row_index}"] = config.display_name
        worksheet[f"B{row_index}"] = specialty.test_result
        worksheet[f"C{row_index}"] = _round_label(specialty.round_no)
        worksheet[f"D{row_index}"] = report_path.name
        row_index += 1


def _write_bug_sheet(worksheet, bug_rows: list[dict[str, str]]) -> None:
    if not bug_rows:
        return

    headers = [cell.value for cell in worksheet[1] if cell.value]
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)
    for row_index, bug_row in enumerate(bug_rows, start=2):
        for column_index, header in enumerate(headers, start=1):
            worksheet.cell(row=row_index, column=column_index, value=bug_row.get(header, ""))


def write_summary_workbook(
    summary_path: Path,
    project: ProjectInfo,
    specialties: list[tuple[SpecialtyExecution, SpecialtyConfig, Path]],
    bug_rows: list[dict[str, str]],
) -> None:
    workbook = load_workbook(summary_path)

    if "总览" in workbook.sheetnames:
        overview_sheet = workbook["总览"]
        _write_overview_sheet(overview_sheet, project, specialties)

    for specialty, config, _ in specialties:
        sheet = workbook[config.summary_sheet_name]
        if "总览" in workbook.sheetnames:
            _write_common_fields(sheet, project, specialty, config)
        else:
            _write_sample_summary_sheet(sheet, project, specialty, config)

    bug_sheet = workbook["Bug_MonkeyList"]
    _write_bug_sheet(bug_sheet, bug_rows)
    workbook.save(summary_path)


def _write_sample_summary_sheet(worksheet, project: ProjectInfo, specialty: SpecialtyExecution, config: SpecialtyConfig) -> None:
    title_row = None
    for row_index in range(1, worksheet.max_row + 1):
        cell_value = worksheet[f"A{row_index}"].value
        if isinstance(cell_value, str) and "测试报告" in cell_value:
            title_row = row_index
            break

    if specialty.specialty_code == "stress":
        worksheet["A1"] = specialty.extensions.get("stress_summary") or f"{project.project_name}压力稳定性{numeric_round_text(specialty.round_no)}测试{specialty.test_result.upper()}，此轮发现问题0个，测试详情如下："
        worksheet["A2"] = build_summary_title(project, specialty, config)
        worksheet["B3"] = project.project_name
        worksheet["D3"] = numeric_round_text(specialty.round_no)
        worksheet["B4"] = project.project_stage
        worksheet["D4"] = specialty.test_result
        worksheet["B5"] = project.reporter
        worksheet["D5"] = specialty.test_cycle
        worksheet["D7"] = int(str(specialty.extensions.get("sample_count") or specialty.device_count))
        worksheet["B8"] = project.software_version
        return

    if specialty.specialty_code == "monkey2":
        worksheet["A3"] = chinese_round_text(specialty.round_no)
        worksheet["B3"] = specialty.test_cycle
        worksheet["C3"] = project.software_version
        worksheet["D3"] = specialty.test_result
        worksheet["K3"] = specialty.test_result
        worksheet["A5"] = build_summary_title(project, specialty, config)
        worksheet["D6"] = project.project_name
        worksheet["H6"] = chinese_round_text(specialty.round_no)
        worksheet["B7"] = project.platform
        worksheet["D7"] = project.bom
        worksheet["F7"] = project.rom_ram
        worksheet["H7"] = project.software_version
        worksheet["B8"] = int(str(specialty.device_count))
        worksheet["D8"] = specialty.test_cycle
        worksheet["H8"] = project.reporter
        return

    if title_row is None:
        return

    worksheet[f"A{title_row}"] = build_summary_title(project, specialty, config)
    worksheet[f"B{title_row + 1}"] = _summary_type_label(specialty, config)
    worksheet[f"D{title_row + 1}"] = specialty.round_no
    worksheet[f"B{title_row + 2}"] = project.project_stage
    worksheet[f"D{title_row + 2}"] = specialty.test_result
    worksheet[f"B{title_row + 3}"] = project.reporter
    worksheet[f"D{title_row + 3}"] = specialty.test_cycle
    worksheet[f"B{title_row + 4}"] = specialty.bug_summary
    if specialty.specialty_code in {"mtbf", "power_cycle", "sleep_wakeup", "gpu", "standby", "ddr", "online_ui_adb_monkey"} and title_row + 6 <= worksheet.max_row:
        worksheet[f"C{title_row + 6}"] = project.software_version


def _summary_type_label(specialty: SpecialtyExecution, config: SpecialtyConfig) -> str:
    labels = {
        "mtbf": "MTBF",
        "power_cycle": "开关机压力测试",
        "sleep_wakeup": "休眠唤醒压力测试",
        "gpu": "GPU压力测试",
        "standby": "待机专项测试",
        "ddr": "DDR专项",
        "online_ui_adb_monkey": "联机UI+ADB重启与原生Monkey压力测试",
    }
    return labels.get(specialty.specialty_code, config.display_name)
