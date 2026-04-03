import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


TOOL_ROOT = Path(__file__).resolve().parents[1]
if str(TOOL_ROOT) not in sys.path:
    sys.path.insert(0, str(TOOL_ROOT))

from report_builder.config import get_specialty_configs
from report_builder.template_factory import ensure_builtin_templates


class TemplateFactoryTests(unittest.TestCase):
    def test_builtin_templates_are_bootstrapped(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            templates_root = Path(temp_dir) / "templates"

            template_paths = ensure_builtin_templates(templates_root)

            self.assertTrue(template_paths.summary.exists())
            for specialty in get_specialty_configs().values():
                self.assertTrue((templates_root / specialty.template_relative_path).exists())

            workbook = load_workbook(template_paths.summary)
            self.assertIn("总览", workbook.sheetnames)
            self.assertIn("MTBF", workbook.sheetnames)
            self.assertIn("Bug_MonkeyList", workbook.sheetnames)


if __name__ == "__main__":
    unittest.main()
