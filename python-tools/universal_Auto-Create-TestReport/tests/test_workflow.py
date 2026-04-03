import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


TOOL_ROOT = Path(__file__).resolve().parents[1]
if str(TOOL_ROOT) not in sys.path:
    sys.path.insert(0, str(TOOL_ROOT))

from report_builder.workflow import run_generation


class WorkflowTests(unittest.TestCase):
    def test_generation_uses_auto_title_when_subreport_title_is_empty(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            input_path = temp_path / "report_input.xlsx"
            output_root = temp_path / "output"

            workbook = Workbook()
            project_sheet = workbook.active
            project_sheet.title = "项目基础信息"
            project_sheet.append([
                "project_name",
                "project_stage",
                "brand",
                "platform",
                "bom",
                "rom_ram",
                "software_version",
                "reporter",
                "report_date",
                "output_root",
                "template_profile",
            ])
            project_sheet.append([
                "KO5",
                "Beta",
                "TECNO",
                "MT6789",
                "P1",
                "256G+8G",
                "KO5-16.2.0.116(OP001PF001AZ)_SU",
                "李小红",
                "2026-03-17",
                str(output_root),
                "default",
            ])

            execution_sheet = workbook.create_sheet("专项执行清单")
            execution_sheet.append([
                "specialty_code",
                "specialty_name",
                "enabled",
                "round_no",
                "test_cycle",
                "test_result",
                "device_count",
                "bug_summary",
                "subreport_title",
            ])
            execution_sheet.append([
                "power_cycle",
                "开关机",
                "TRUE",
                5,
                "2026/3/10-2026/3/14",
                "PASS",
                10,
                "Blocker：0、Critical：0、Major：0",
                "",
            ])

            extension_sheet = workbook.create_sheet("专项扩展字段")
            extension_sheet.append(["specialty_code", "field_key", "field_value"])

            bug_sheet = workbook.create_sheet("Bug汇总导入")
            bug_sheet.append([
                "Issue key",
                "Project",
                "Affects Version/s",
                "Component/s",
                "Priority",
                "Custom field (Risk)",
                "Summary",
                "Description",
                "Reporter",
            ])
            workbook.save(input_path)

            result = run_generation(input_path=input_path)

            self.assertEqual(len(result.failed_specialties), 0)
            self.assertEqual(len(result.generated_reports), 1)
            report_path = result.generated_reports[0]
            self.assertIn("16.2.0.116(OP001PF001AZ)_SU版本开关机专项第五轮测试报告----PASS.xlsx", report_path.name)

    def test_generation_creates_summary_and_enabled_specialties(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            input_path = temp_path / "report_input.xlsx"
            output_root = temp_path / "output"

            workbook = Workbook()
            project_sheet = workbook.active
            project_sheet.title = "项目基础信息"
            project_sheet.append([
                "project_name",
                "project_stage",
                "brand",
                "platform",
                "bom",
                "rom_ram",
                "software_version",
                "reporter",
                "report_date",
                "output_root",
                "template_profile",
            ])
            project_sheet.append([
                "KO5",
                "Beta",
                "TECNO",
                "MT6789",
                "P1",
                "256G+8G",
                "KO5-16.2.0.116(OP001PF001AZ)_SU",
                "李小红",
                "2026-03-17",
                str(output_root),
                "default",
            ])

            execution_sheet = workbook.create_sheet("专项执行清单")
            execution_sheet.append([
                "specialty_code",
                "specialty_name",
                "enabled",
                "round_no",
                "test_cycle",
                "test_result",
                "device_count",
                "bug_summary",
                "subreport_title",
            ])
            execution_sheet.append([
                "mtbf",
                "MTBF",
                "TRUE",
                5,
                "2026/3/10-2026/3/17",
                "PASS",
                10,
                "Blocker：0、Critical：0、Major：0",
                "KO5项目Beta阶段MTBF专项第五轮测试报告",
            ])
            execution_sheet.append([
                "stress",
                "Stress",
                "TRUE",
                3,
                "2026/3/10-2026/3/17",
                "PASS",
                20,
                "Blocker：0、Critical：0、Major：0",
                "KO5项目Beta阶段Stress专项第三轮测试报告",
            ])
            execution_sheet.append([
                "power_cycle",
                "开关机",
                "TRUE",
                5,
                "2026/3/10-2026/3/14",
                "PASS",
                10,
                "Blocker：0、Critical：0、Major：0",
                "KO5项目Beta阶段开关机专项第五轮测试报告",
            ])
            execution_sheet.append([
                "sleep_wakeup",
                "休眠唤醒",
                "TRUE",
                5,
                "2026/3/10-2026/3/17",
                "PASS",
                10,
                "Blocker：0、Critical：0、Major：0",
                "KO5项目Beta阶段休眠唤醒专项第五轮测试报告",
            ])
            execution_sheet.append([
                "gpu",
                "GPU",
                "TRUE",
                4,
                "2026/2/25-2026/3/4",
                "PASS",
                5,
                "Blocker：0、Critical：0、Major：0",
                "KO5项目Beta阶段GPU专项第四轮测试报告",
            ])
            execution_sheet.append([
                "standby",
                "待机",
                "TRUE",
                3,
                "2026/2/25-2026/3/4",
                "PASS",
                10,
                "Blocker：0、Critical：0、Major：0",
                "KO5项目Beta阶段待机专项第三轮测试报告",
            ])
            execution_sheet.append([
                "ddr",
                "DDR",
                "TRUE",
                1,
                "2026/1/20-2026/1/27",
                "PASS",
                50,
                "Blocker：0、Critical：0、Major：0",
                "KO5项目Beta阶段DDR专项第一轮测试报告",
            ])
            execution_sheet.append([
                "online_ui_adb_monkey",
                "联机UI+ADB重启与原生Monkey",
                "TRUE",
                3,
                "2026/3/10-2026/3/19",
                "PASS",
                10,
                "Blocker：0、Critical：0、Major：0",
                "KO5项目Beta阶段联机UI+ADB重启+恢复出厂开关机专项第三轮测试报告",
            ])
            execution_sheet.append([
                "monkey2",
                "Monkey2.0",
                "TRUE",
                6,
                "2026/2/26-2026/3/7",
                "FAIL",
                120,
                "Blocker：5、Critical：0、Major：0",
                "KO5项目Beta阶段Monkey专项第六轮测试报告",
            ])

            extension_sheet = workbook.create_sheet("专项扩展字段")
            extension_sheet.append(["specialty_code", "field_key", "field_value"])
            extension_sheet.append(["stress", "stress_summary", "KO5压力稳定性第3轮测试PASS，此轮发现问题0个"])
            extension_sheet.append(["stress", "sample_count", "20"])

            bug_sheet = workbook.create_sheet("Bug汇总导入")
            bug_sheet.append([
                "Issue key",
                "Project",
                "Affects Version/s",
                "Component/s",
                "Priority",
                "Custom field (Risk)",
                "Summary",
                "Description",
                "Reporter",
            ])
            bug_sheet.append([
                "KO5OS16AEE-748",
                "KO5",
                "KO5-16.2.0.116(OP001PF001AZ)_SU",
                "Reboot",
                "Blocker",
                "occasional",
                "[MonkeyAEE] SWT system_server",
                "示例问题",
                "xiaohongli.tinno",
            ])
            workbook.save(input_path)

            result = run_generation(input_path=input_path)

            self.assertEqual(len(result.generated_reports), 9)
            self.assertEqual(len(result.failed_specialties), 0)
            self.assertTrue(result.summary_path.exists())
            self.assertTrue((result.output_dir / "压力测试").exists())
            self.assertTrue((result.output_dir / "monkey").exists())

            summary_book = load_workbook(result.summary_path)
            self.assertNotIn("总览", summary_book.sheetnames)
            self.assertEqual(summary_book["MTBF"]["A1"].value, "测试轮数")
            self.assertEqual(summary_book["MTBF"]["A9"].value, "KO5项目Beta阶段16.2.0.116(OP001PF001AZ)_SU版本MTBF专项第五轮测试报告")
            self.assertEqual(summary_book["Stress"]["A2"].value, "KO5项目Beta阶段稳定性Stress测试报告【PASS】")
            self.assertEqual(summary_book["开关机"]["A8"].value, "KO5项目Beta阶段16.2.0.116(OP001PF001AZ)_SU版本开关机专项第五轮测试报告")
            self.assertEqual(summary_book["休眠唤醒"]["A8"].value, "KO5项目Beta阶段16.2.0.116(OP001PF001AZ)_SU版本休眠唤醒专项第五轮测试报告")
            self.assertEqual(summary_book["GPU"]["A7"].value, "KO5项目Beta阶段16.2.0.116(OP001PF001AZ)_SU版本GPU专项第四轮测试报告")
            self.assertEqual(summary_book["待机"]["A6"].value, "KO5项目Beta阶段16.2.0.116(OP001PF001AZ)_SU版本待机专项第三轮测试报告")
            self.assertEqual(summary_book["DDR"]["A4"].value, "KO5项目Beta阶段16.2.0.116(OP001PF001AZ)_SU版本DDR专项第一轮测试报告")
            self.assertEqual(summary_book["联机UI重启+ADB重启与原生Monkey"]["A6"].value, "KO5项目Beta阶段16.2.0.116(OP001PF001AZ)_SU版本联机UI+ADB重启+恢复出厂开关机专项第三轮测试报告")
            self.assertEqual(summary_book["Monkey2.0"]["A5"].value, "KO5项目Beta阶段16.2.0.116(OP001PF001AZ)_SU版本(第六轮)monkey专项测试报告")
            self.assertEqual(summary_book["Bug_MonkeyList"]["A2"].value, "KO5OS16AEE-748")

            mtbf_report = next(path for path in result.generated_reports if "MTBF" in path.name)
            stress_report = next(path for path in result.generated_reports if "Stress" in path.name)
            power_cycle_report = next(path for path in result.generated_reports if "开关机" in path.name)
            sleep_report = next(path for path in result.generated_reports if "休眠唤醒" in path.name)
            gpu_report = next(path for path in result.generated_reports if "GPU专项第四轮测试报告" in path.name)
            standby_report = next(path for path in result.generated_reports if "待机专项第三轮测试报告" in path.name)
            ddr_report = next(path for path in result.generated_reports if "DDR专项第一轮测试报告" in path.name)
            online_report = next(path for path in result.generated_reports if "联机UI+ADB重启+恢复出厂开关机专项第三轮测试报告" in path.name)
            monkey_report = next(path for path in result.generated_reports if "monkey专项测试报告" in path.name)
            self.assertTrue(mtbf_report.exists())
            self.assertTrue(stress_report.exists())
            self.assertTrue(power_cycle_report.exists())
            self.assertTrue(sleep_report.exists())
            self.assertTrue(gpu_report.exists())
            self.assertTrue(standby_report.exists())
            self.assertTrue(ddr_report.exists())
            self.assertTrue(online_report.exists())
            self.assertTrue(monkey_report.exists())
            self.assertIn("16.2.0.116(OP001PF001AZ)_SU版本开关机专项第五轮测试报告----PASS.xlsx", power_cycle_report.name)
            self.assertIn("16.2.0.116(OP001PF001AZ)_SU版本休眠唤醒专项第五轮测试报告----PASS.xlsx", sleep_report.name)
            self.assertIn("16.2.0.116(OP001PF001AZ)_SU版本GPU专项第四轮测试报告----PASS.xlsx", gpu_report.name)
            self.assertIn("16.2.0.116(OP001PF001AZ)_SU版本待机专项第三轮测试报告----PASS.xlsx", standby_report.name)
            self.assertIn("16.2.0.116(OP001PF001AZ)_SU版本DDR专项第一轮测试报告----PASS.xlsx", ddr_report.name)
            self.assertIn("16.2.0.116(OP001PF001AZ)_SU版本联机UI+ADB重启+恢复出厂开关机专项第三轮测试报告----PASS.xlsx", online_report.name)
            self.assertIn("16.2.0.116(OP001PF001AZ)_SU版本(第六轮)monkey专项测试报告----FAIL.xlsx", monkey_report.name)
            self.assertEqual(ddr_report.parent.name, "DDR")

            stress_book = load_workbook(stress_report)
            self.assertEqual(stress_book.sheetnames[:2], ["测试报告", "测试结果"])
            self.assertEqual(
                stress_book["测试报告"]["B2"].value.split("\n")[0],
                "KO5项目Beta阶段16.2.0.116(OP001PF001AZ)_SU版本稳定性专项内存泄露（Stress）第三轮测试报告",
            )
            self.assertEqual(stress_book["测试报告"]["C6"].value, 20)
            self.assertIn("KO5压力稳定性第3轮测试PASS", stress_book["测试报告"]["B12"].value)

            power_cycle_book = load_workbook(power_cycle_report)
            self.assertEqual(
                power_cycle_book["测试报告"]["B2"].value.split("\n")[0],
                "KO5项目Beta阶段16.2.0.116(OP001PF001AZ)_SU版本开关机专项第五轮测试报告",
            )

            sleep_book = load_workbook(sleep_report)
            self.assertEqual(
                sleep_book["测试报告"]["B2"].value.split("\n")[0],
                "KO5项目Beta阶段16.2.0.116(OP001PF001AZ)_SU版本休眠唤醒专项第五轮测试报告",
            )

            gpu_book = load_workbook(gpu_report)
            self.assertEqual(
                gpu_book["测试报告"]["B2"].value.split("\n")[0],
                "KO5项目Beta阶段16.2.0.116(OP001PF001AZ)_SU版本GPU专项第四轮测试报告",
            )

            standby_book = load_workbook(standby_report)
            self.assertEqual(
                standby_book["测试报告"]["B2"].value.split("\n")[0],
                "KO5项目Beta阶段16.2.0.116(OP001PF001AZ)_SU版本待机专项第三轮测试报告",
            )

            ddr_book = load_workbook(ddr_report)
            self.assertEqual(
                ddr_book["测试报告"]["B2"].value.split("\n")[0],
                "KO5项目Beta阶段16.2.0.116(OP001PF001AZ)_SU版本DDR专项第一轮测试报告",
            )

            online_book = load_workbook(online_report)
            self.assertEqual(
                online_book["测试报告"]["B2"].value.split("\n")[0],
                "KO5项目Beta阶段16.2.0.116(OP001PF001AZ)_SU版本联机UI+ADB重启+恢复出厂开关机专项第三轮测试报告",
            )

            monkey_book = load_workbook(monkey_report)
            self.assertEqual(monkey_book["result"]["A1"].value, "KO5项目Beta阶段16.2.0.116(OP001PF001AZ)_SU版本(第六轮)monkey专项测试报告")
            self.assertEqual(monkey_book["result"]["D4"].value, "KO5")
            self.assertEqual(monkey_book["result"]["F4"].value, "Beta")
            self.assertEqual(monkey_book["result"]["B6"].value, "KO5-16.2.0.116(OP001PF001AZ)_SU")


if __name__ == "__main__":
    unittest.main()
