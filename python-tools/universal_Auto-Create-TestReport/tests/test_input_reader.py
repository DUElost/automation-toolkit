import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


TOOL_ROOT = Path(__file__).resolve().parents[1]
if str(TOOL_ROOT) not in sys.path:
    sys.path.insert(0, str(TOOL_ROOT))

from report_builder.input_reader import InputValidationError, read_input_workbook


class InputReaderTests(unittest.TestCase):
    def test_missing_required_sheet_raises_validation_error(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "input.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "项目基础信息"
            sheet.append(["project_name", "project_stage", "software_version", "reporter", "report_date", "output_root", "template_profile"])
            sheet.append(["KO5", "Beta", "KO5-16.2.0.116", "李小红", "2026-03-17", str(Path(temp_dir) / "output"), "default"])
            workbook.save(workbook_path)

            with self.assertRaises(InputValidationError) as context:
                read_input_workbook(workbook_path)

            self.assertIn("专项执行清单", str(context.exception))

    def test_extension_sheet_is_grouped_by_specialty(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "input.xlsx"
            workbook = Workbook()

            project_sheet = workbook.active
            project_sheet.title = "项目基础信息"
            project_sheet.append([
                "project_name",
                "project_stage",
                "software_version",
                "reporter",
                "report_date",
                "output_root",
                "template_profile",
            ])
            project_sheet.append([
                "KO5",
                "Beta",
                "KO5-16.2.0.116",
                "李小红",
                "2026-03-17",
                str(Path(temp_dir) / "output"),
                "default",
            ])

            execution_sheet = workbook.create_sheet("专项执行清单")
            execution_sheet.append([
                "specialty_code",
                "specialty_name",
                "enabled",
                "round_no",
                "test_cycle",
                "test_result",
                "device_count",
                "bug_summary",
                "subreport_title",
            ])
            execution_sheet.append([
                "stress",
                "Stress",
                "TRUE",
                3,
                "2026/3/10-2026/3/17",
                "PASS",
                20,
                "Blocker：0、Critical：0、Major：0",
                "KO5 Stress 第3轮测试报告",
            ])

            extension_sheet = workbook.create_sheet("专项扩展字段")
            extension_sheet.append(["specialty_code", "field_key", "field_value"])
            extension_sheet.append(["stress", "stress_summary", "KO5压力稳定性第3轮测试PASS"])
            extension_sheet.append(["stress", "sample_count", "20"])

            workbook.save(workbook_path)

            report_input = read_input_workbook(workbook_path)

            self.assertEqual(report_input.project.project_name, "KO5")
            self.assertEqual(len(report_input.specialties), 1)
            self.assertEqual(report_input.specialties[0].extensions["stress_summary"], "KO5压力稳定性第3轮测试PASS")
            self.assertEqual(report_input.specialties[0].extensions["sample_count"], "20")


if __name__ == "__main__":
    unittest.main()
